"""
Tests de l'API (suivi, génération en file, validation humaine), sans
infrastructure : base SQLite temporaire, générateur factice à la place de
Mistral, génération exécutée immédiatement au lieu d'une tâche de fond.

    python -m unittest discover -s tests
"""
import os
import shutil
import uuid
import sys
import tempfile
import unittest
from pathlib import Path

TMP = Path(tempfile.mkdtemp(prefix="conformite-tests-"))
os.environ["DATABASE_URL"] = f"sqlite+pysqlite:///{TMP / 'test.db'}"
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from docx import Document as Docx  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import api  # noqa: E402
import db  # noqa: E402
import extraction  # noqa: E402
import generate  # noqa: E402
import taches  # noqa: E402

APPELS_GENERATION: list[str] = []
CONTEXTES: list[dict] = []


def faux_generer_document(_generateur, doc, produit, pays_origine, dossier_sortie, **contexte):
    """Remplace RAG + Mistral : écrit un vrai DOCX minimal, et refuse comme le vrai une pièce à fournir."""
    if not doc.a_rediger:
        raise generate.GenerationImpossible("pièce à fournir")
    APPELS_GENERATION.append(doc.id)
    CONTEXTES.append(contexte)
    dossier_sortie.mkdir(parents=True, exist_ok=True)
    chemin = dossier_sortie / f"{doc.id}.docx"
    d = Docx()
    d.add_paragraph(f"Projet : {doc.nom} — {produit}")
    d.save(chemin)
    return {"fichier": str(chemin), "sources": [{"texte_source": "Arrêtés 2853-2856", "date_version": "2015-08-04"}]}


def faux_extraire(chemin, champs, piece):
    """Remplace OCR + Mistral : renvoie un résultat au format réel."""
    if chemin.read_bytes().startswith(b"illisible"):
        raise extraction.ExtractionImpossible("Aucun texte lisible dans le document.")
    resultats = [{"nom": c["nom"], "libelle": c["libelle"], "type": c["type"], "valeur": "X",
                  "valeur_normalisee": "X", "citation": "X", "verification": extraction.VERIFIE} for c in champs]
    return {"champs": resultats, "resume": {extraction.VERIFIE: len(champs), extraction.CITATION_INTROUVABLE: 0,
            extraction.VALEUR_HORS_CITATION: 0, extraction.ABSENT: 0}, "caracteres_lus": 10, "texte": "texte lu",
            "modele": "faux"}


class TestApi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        generate.OUTPUT_DIR = TMP / "output"
        api.OUTPUT_DIR = (TMP / "output").resolve()
        generate.generer_document = faux_generer_document
        generate.PROFIL_ENTREPRISE = TMP / "profil_absent.yaml"  # jamais le profil réel de l'entreprise
        taches._obtenir_generateur = lambda: None
        taches.soumettre = taches.executer_generation  # exécution immédiate
        taches.soumettre_extraction = taches.executer_extraction
        cls._vrai_extraire = extraction.extraire  # rétabli à la fin : les autres tests lisent pour de vrai
        extraction.extraire = faux_extraire
        import memoire
        memoire.RACINE_DEFAUT = TMP / "dossiers_valides_absents"  # jamais les dossiers acceptés réels
        memoire.DOSSIER_MEMOIRE = TMP / "memoire_absente"
        cls.client = TestClient(api.app)
        cls.client.__enter__()  # déclenche le lifespan (création des tables)

    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None, None, None)
        extraction.extraire = cls._vrai_extraire
        shutil.rmtree(TMP, ignore_errors=True)

    def setUp(self):
        APPELS_GENERATION.clear()
        api.API_KEY = None

    def creer(self, pays="chine", classe="IIB", fournisseur=None, produit="Prothèse orthopédique de hanche"):
        # fournisseur distinct par dossier : la bibliothèque fournisseur ne relie pas les tests entre eux
        r = self.client.post("/dossiers", json={
            "pays_origine": pays, "produit": produit,
            "classe": classe, "fournisseur": fournisseur or f"Fournisseur{uuid.uuid4().hex[:8]}", "cree_par": "Testeur"})
        self.assertEqual(r.status_code, 201, r.text)
        return r.json()

    @staticmethod
    def piece(dossier, code):
        return next(d for d in dossier["documents"] if d["code"] == code)

    def test_creation_fige_la_decision_du_moteur_de_regles(self):
        d = self.creer()
        # règles 2.5, Chine IIb : socle + ISO + NMPA/FSC + échantillon (valeur non saisie) + évaluation clinique
        # (hors UE, loi 84-12 art. 15) + déclaration de l'établissement (loi 84-12 art. 7)
        self.assertEqual(d["compteurs"]["total"], 15)
        natures = {x["code"]: (x["nature"], x["statut"]) for x in d["documents"]}
        self.assertEqual(natures["demande_signee"], ("a_rediger", "a_generer"))
        self.assertEqual(natures["piece_specifique_chine"], ("a_fournir", "a_obtenir"))
        self.assertEqual(natures["iso_13485"], ("a_fournir", "a_obtenir"))
        self.assertEqual(d["statut"], "en_preparation")
        self.assertEqual(len(d["regles_version"]), 12)
        self.assertEqual(d["evenements"][0]["action"], "dossier_cree")

    def test_classe_I_avec_iso_et_conditions(self):
        d = self.creer(pays="union_europeenne", classe="I")
        self.assertIn("iso_13485", [x["code"] for x in d["documents"]])  # règles 2.4 : toutes les classes
        r = self.client.post("/dossiers", json={"pays_origine": "chine", "produit": "Moniteur", "classe": "IIB",
                                                "equipement": True, "valeur_unitaire_usd": 4000, "cree_par": "Testeur"})
        codes = [x["code"] for x in r.json()["documents"]]
        self.assertTrue(r.json()["equipement"])
        self.assertEqual(r.json()["valeur_unitaire_usd"], 4000)
        self.assertIn("facture_proforma", codes)
        self.assertIn("note_descriptive", codes)
        self.assertNotIn("echantillon_modele_vente", codes)

    def test_numero_remarque_et_source_exposes(self):
        d = self.creer(pays="union_europeenne", classe="IIB")
        ce = self.piece(d, "piece_specifique_union_europeenne")
        self.assertEqual(ce["numero"], 4)
        self.assertIn("2023/607", ce["remarque"])
        self.assertIn("2855-15", self.piece(d, "demande_signee")["source"])
        self.assertEqual([x["numero"] for x in d["documents"]][:3], [1, 2, 3])  # ordre du dossier déposé

    def test_sous_classes_is_im_acceptees(self):
        d = self.creer(pays="union_europeenne", classe="IS")
        codes = [x["code"] for x in d["documents"]]
        self.assertIn("piece_specifique_union_europeenne", codes)
        self.assertIn("iso_13485", codes)  # règles 2.4 : toutes les classes
        d = self.creer(pays="union_europeenne", classe="IR")  # classe I réutilisable (MDR)
        self.assertEqual(d["classe"], "IR")
        self.assertIn("piece_specifique_union_europeenne", [x["code"] for x in d["documents"]])

    def test_correspondances_et_nouvelles_origines(self):
        r = self.client.get("/correspondances")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(r.json()["pays"]), 7)
        pays = self.client.get("/pays").json()
        self.assertEqual(pays["coree_du_sud"]["statut"], "verifie")
        self.assertEqual(pays["maroc"]["statut"], "verifie")
        d = self.creer(pays="etats_unis")
        self.assertIn("piece_specifique_etats_unis", [x["code"] for x in d["documents"]])
        self.assertEqual(self.client.get("/correspondances/comparer/inconnu").status_code, 404)

    def test_generation_ne_redige_que_les_pieces_a_rediger(self):
        d = self.creer()
        r = self.client.post(f"/dossiers/{d['id']}/generer", json={"acteur": "Testeur"})
        self.assertEqual(r.status_code, 202, r.text)
        # Mistral ne rédige que la lettre ; les deux formulaires sont remplis par le code
        self.assertEqual(APPELS_GENERATION, ["demande_signee"])
        for code in ("fiche_signaletique", "certificat_enregistrement_annexe2"):
            piece = self.piece(r.json(), code)
            self.assertTrue(piece["fichier_disponible"], code)
            self.assertIn("rempli par le code", piece["sources"][0]["texte_source"])
        d = self.client.get(f"/dossiers/{d['id']}").json()
        for x in d["documents"]:
            attendu = "a_valider" if x["nature"] == "a_rediger" else "a_obtenir"
            self.assertEqual(x["statut"], attendu, x["code"])
        self.assertTrue(self.piece(d, "demande_signee")["fichier_disponible"])

    def test_contexte_transmis_a_la_redaction(self):
        d = self.creer()
        CONTEXTES.clear()
        self.client.post(f"/dossiers/{d['id']}/generer", json={"acteur": "Testeur"})
        ctx = CONTEXTES[0]
        self.assertEqual(ctx["classe"], "IIB")
        self.assertEqual(ctx["fournisseur"], d["fournisseur"])
        # la liste des pièces transmise à Mistral est exactement celle du moteur de règles
        self.assertEqual(ctx["pieces_du_dossier"], [x["nom"] for x in d["documents"]])

    def test_piece_a_fournir_jamais_generee(self):
        d = self.creer()
        nmpa = self.piece(d, "piece_specifique_chine")
        r = self.client.post(f"/dossiers/{d['id']}/documents/{nmpa['id']}/generer", json={"acteur": "Testeur"})
        self.assertEqual(r.status_code, 409)
        self.assertIn("jamais rédigée", r.json()["detail"])
        self.assertEqual(APPELS_GENERATION, [])

    def test_cycle_complet_de_validation(self):
        d = self.creer(pays="union_europeenne", classe="IIA")
        self.client.post(f"/dossiers/{d['id']}/generer", json={"acteur": "Testeur"})
        d = self.client.get(f"/dossiers/{d['id']}").json()
        for x in d["documents"]:
            url = f"/dossiers/{d['id']}/documents/{x['id']}/valider"
            if x["nature"] == "a_fournir":
                r = self.client.post(url, json={"validateur": "Relecteur"})
                self.assertEqual(r.status_code, 422, "commentaire obligatoire pour une pièce à fournir")
                r = self.client.post(url, json={"validateur": "Relecteur", "commentaire": "Reçu le 30/09, réf. X"})
            else:
                r = self.client.post(url, json={"validateur": "Relecteur"})
            self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertEqual(d["statut"], "pret_pour_depot_manuel")
        self.assertIsNotNone(d["prochain_creneau_depot"])
        self.assertTrue(all(x["valide_par"] == "Relecteur" for x in d["documents"]))

    def test_rejet_motif_obligatoire_puis_regeneration(self):
        d = self.creer()
        self.client.post(f"/dossiers/{d['id']}/generer", json={"acteur": "Testeur"})
        demande = self.piece(self.client.get(f"/dossiers/{d['id']}").json(), "demande_signee")
        url = f"/dossiers/{d['id']}/documents/{demande['id']}"
        self.assertEqual(self.client.post(f"{url}/rejeter", json={"validateur": "Relecteur"}).status_code, 422)
        r = self.client.post(f"{url}/rejeter", json={"validateur": "Relecteur", "commentaire": "Recopie les sources"})
        self.assertEqual(self.piece(r.json(), "demande_signee")["statut"], "rejete")
        APPELS_GENERATION.clear()
        r = self.client.post(f"{url}/generer", json={"acteur": "Testeur"})
        self.assertEqual(r.status_code, 202)
        self.assertEqual(APPELS_GENERATION, ["demande_signee"])
        self.assertEqual(self.piece(r.json(), "demande_signee")["statut"], "a_valider")

    def test_validation_impossible_avant_redaction(self):
        d = self.creer()
        demande = self.piece(d, "demande_signee")
        r = self.client.post(f"/dossiers/{d['id']}/documents/{demande['id']}/valider", json={"validateur": "Relecteur"})
        self.assertEqual(r.status_code, 409)

    def test_telechargement_du_projet(self):
        d = self.creer()
        self.client.post(f"/dossiers/{d['id']}/generer", json={"acteur": "Testeur"})
        demande = self.piece(d, "demande_signee")
        r = self.client.get(f"/dossiers/{d['id']}/documents/{demande['id']}/fichier")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"PK"))  # un .docx est une archive zip

    def test_apercu_du_projet(self):
        d = self.creer()
        self.client.post(f"/dossiers/{d['id']}/generer", json={"acteur": "Testeur"})
        demande = self.piece(d, "demande_signee")
        r = self.client.get(f"/dossiers/{d['id']}/documents/{demande['id']}/apercu")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("Projet : Lettre de demande", r.json()["paragraphes"][0]["texte"])
        nmpa = self.piece(d, "piece_specifique_chine")
        self.assertEqual(self.client.get(f"/dossiers/{d['id']}/documents/{nmpa['id']}/apercu").status_code, 404)

    def deposer(self, d, code, contenu=b"%PDF-1.4 certificat", nom="certificat.pdf"):
        piece = self.piece(d, code)
        return self.client.post(f"/dossiers/{d['id']}/documents/{piece['id']}/document-recu",
                                data={"acteur": "Testeur"}, files={"fichier": (nom, contenu, "application/pdf")})

    def test_lecture_du_document_recu(self):
        d = self.creer()
        self.assertTrue(self.piece(d, "iso_13485")["lisible_par_agent"])
        r = self.deposer(d, "iso_13485")
        self.assertEqual(r.status_code, 202, r.text)
        iso = self.piece(r.json(), "iso_13485")
        self.assertEqual(iso["extraction_statut"], "terminee")
        noms = [c["nom"] for c in iso["extraction"]["champs"]]
        self.assertEqual(noms[:7], ["numero", "emetteur", "titulaire", "produit_couvert", "date_emission", "date_expiration", "norme"])
        self.assertIn("cl_certificat_iso_1", noms)  # + les points « lecture » de la checklist de l'entreprise
        self.assertNotIn("texte", iso["extraction"])  # le texte intégral n'est pas renvoyé au navigateur
        self.assertEqual(iso["statut"], "a_obtenir")  # lire n'est pas valider : la décision reste humaine
        actions = [e["action"] for e in r.json()["evenements"]]
        self.assertIn("document_recu_depose", actions)
        self.assertIn("lecture_terminee", actions)
        orig = self.client.get(f"/dossiers/{d['id']}/documents/{iso['id']}/document-recu")
        self.assertEqual(orig.content, b"%PDF-1.4 certificat")

    def test_donnees_du_dispositif(self):
        d = self.creer(pays="union_europeenne")
        url = f"/dossiers/{d['id']}/donnees-dispositif"
        r = self.client.get(url).json()
        champs = {c["id"]: c for s in r["sections"] for c in s["champs"]}
        self.assertEqual(champs["designation"]["valeur"], "Prothèse orthopédique de hanche")
        self.assertEqual(champs["indications"]["provenance"], "manquant")
        self.assertGreater(r["a_completer"], 5)
        r = self.client.put(url, json={"acteur": "Testeur", "valeurs": {"indications": "Arthroplastie", "categorie": ""}})
        champs = {c["id"]: c for s in r.json()["sections"] for c in s["champs"]}
        self.assertEqual((champs["indications"]["valeur"], champs["indications"]["provenance"]), ("Arthroplastie", "saisie"))
        self.assertEqual(self.client.put(url, json={"acteur": "Testeur", "valeurs": {"pirate": "x"}}).status_code, 422)
        journal = [e["action"] for e in self.client.get(f"/dossiers/{d['id']}").json()["evenements"]]
        self.assertIn("donnees_saisies", journal)
        self.assertIn("formulaire_mis_a_jour", journal)  # fiche et annexe II remises à jour automatiquement
        # la saisie est reprise dans le formulaire rempli par le code
        self.client.post(f"/dossiers/{d['id']}/generer", json={"acteur": "Testeur"})
        fiche = self.piece(self.client.get(f"/dossiers/{d['id']}").json(), "fiche_signaletique")
        apercu = self.client.get(f"/dossiers/{d['id']}/documents/{fiche['id']}/apercu").json()["paragraphes"]
        self.assertIn({"genre": "ligne", "texte": "Indications | Arthroplastie"}, apercu)
        # reprise dans un autre dossier, sans écraser ce qui y est déjà saisi
        autre = self.creer(pays="union_europeenne")
        self.client.put(f"/dossiers/{autre['id']}/donnees-dispositif",
                        json={"acteur": "Testeur", "valeurs": {"indications": "Autre"}})
        r = self.client.post(f"/dossiers/{autre['id']}/donnees-dispositif/reprendre",
                             json={"acteur": "Testeur", "depuis": d["id"]})
        champs = {c["id"]: c for s in r.json()["sections"] for c in s["champs"]}
        self.assertEqual(champs["indications"]["valeur"], "Autre")

    def test_depot_groupe_range_par_le_code(self):
        d = self.creer(pays="union_europeenne")
        fichiers = [("fichiers", (nom, b"%PDF-1.4 " + nom.encode(), "application/pdf")) for nom in
                    ("ISO13485.pdf", "DoC for Morocco.pdf", "facture.pdf", "9-9- Photo.pdf")]
        r = self.client.post(f"/dossiers/{d['id']}/documents-recus", data={"acteur": "Testeur"}, files=fichiers)
        self.assertEqual(r.status_code, 202, r.text)
        affectations = {a["fichier"]: a["piece"] for a in r.json()["affectations"]}
        self.assertEqual(affectations, {"ISO13485.pdf": 5, "DoC for Morocco.pdf": 6, "facture.pdf": None,
                                        "9-9- Photo.pdf": 9})  # photos : reçues (rien à lire)
        dossier = r.json()["dossier"]
        self.assertEqual(self.piece(dossier, "iso_13485")["extraction_statut"], "terminee")
        doc = self.piece(dossier, "declaration_conformite")
        self.assertIn("adresse_fabricant", [c["nom"] for c in doc["extraction"]["champs"]])  # champs des règles actuelles
        self.assertEqual(self.piece(dossier, "notice")["extraction_statut"], None)

    def test_controles_de_la_checklist(self):
        d = self.creer(pays="union_europeenne")
        r = self.client.get(f"/dossiers/{d['id']}/controles")
        self.assertEqual(r.status_code, 200, r.text)
        ids = {x["id"] for x in r.json()["documents"]}
        self.assertIn("declaration_de_conformite", ids)
        self.assertGreater(r.json()["resume"]["en_attente"], 0)
        self.assertIn("Prothèse orthopédique de hanche", r.json()["relance"])
        self.deposer(d, "iso_13485")  # lu par le faux agent : tous les points « vérifiés »
        r = self.client.get(f"/dossiers/{d['id']}/controles").json()
        iso = next(x for x in r["documents"] if x["id"] == "certificat_iso")
        self.assertEqual(iso["elements"][1]["statut"], "ok")  # périmètre du certificat
        sig = "cl_declaration_de_conformite_19"
        r = self.client.post(f"/dossiers/{d['id']}/controles/humain", json={"acteur": "Testeur", "element": sig, "fait": True})
        doc = next(x for x in r.json()["documents"] if x["id"] == "declaration_de_conformite")
        self.assertEqual(doc["elements"][19]["statut"], "humain_fait")
        self.assertEqual(self.client.post(f"/dossiers/{d['id']}/controles/humain",
                                          json={"acteur": "Testeur", "element": "cl_certificat_iso_1", "fait": True}).status_code, 422)
        journal = self.client.get(f"/dossiers/{d['id']}").json()["evenements"]
        self.assertEqual(journal[0]["action"], "point_verifie")

    def test_enchainement_automatique(self):
        # les deux cas de la direction : même si le fournisseur n'a pas tout envoyé, l'agent prépare
        # nos papiers dès que les documents reçus sont lus (les pièces manquantes sont réclamées)
        d = self.creer(pays="union_europeenne")
        APPELS_GENERATION.clear()
        self.deposer(d, "iso_13485")  # seul document reçu
        dossier = self.client.get(f"/dossiers/{d['id']}").json()
        self.assertEqual(self.piece(dossier, "fiche_signaletique")["statut"], "a_valider")  # rempli dès la 1re lecture
        self.assertEqual(APPELS_GENERATION, ["demande_signee"])  # lettre lancée sans attendre les pièces manquantes
        self.assertEqual(self.piece(dossier, "demande_signee")["statut"], "a_valider")
        self.assertEqual(self.piece(dossier, "demande_signee")["valide_par"], None)  # rien n'est validé automatiquement
        self.deposer(d, "notice")  # document arrivé plus tard : la lettre déjà rédigée n'est pas relancée
        self.assertEqual(APPELS_GENERATION, ["demande_signee"])

    def test_bilan_trois_cas(self):
        # cas « incomplet » : seul l'ISO reçu, rien dans la base
        d = self.creer(pays="union_europeenne", fournisseur="Fabricant Bilan", produit="Pansement Bilan")
        self.deposer(d, "iso_13485")
        b = self.client.get(f"/dossiers/{d['id']}/bilan").json()
        self.assertEqual(b["cas"], "incomplet")
        self.assertIn("notice", [l["code"] for l in b["manquants"]])
        self.assertIn("AMMPS", b["relance"])  # relance prête (jamais envoyée)
        notre = {l["code"]: l for l in b["notre_part"]}
        self.assertEqual(notre["demande_signee"]["etat"], "redige")  # nos papiers préparés quand même
        self.assertEqual(notre["quittance_droits"]["etat"], "a_fournir_par_nous")
        # cas « complété par l'agent » : même fabricant, même produit, le fournisseur n'envoie que les certificats
        lisibles = [x["code"] for x in d["documents"] if x["lisible_par_agent"] and x["code"] != "iso_13485"]
        for code in lisibles:
            self.deposer(d, code)
        second = self.creer(pays="union_europeenne", fournisseur="Fabricant Bilan", produit="Pansement Bilan")
        for code in ("iso_13485", "piece_specifique_union_europeenne"):
            self.deposer(second, code)
        b = self.client.get(f"/dossiers/{second['id']}/bilan").json()
        self.assertIn(b["cas"], ("complete_par_agent", "incomplet"))
        self.assertTrue(b["complementaires"])  # pièces complémentaires fournies par l'agent
        self.assertTrue(all(l["par"] == "agent (base)" for l in b["complementaires"]))
        etiquettes = next(l for l in b["notre_part"] if l["code"] == "etiquetage")
        self.assertEqual(etiquettes["etat"], "a_signer")  # étiquettes : à signer et cacheter par nous
        # cas « complet » : le fournisseur envoie tout lui-même
        troisieme = self.creer(pays="union_europeenne", fournisseur="Fabricant Complet", produit="Seringue Complet")
        for code in [x["code"] for x in troisieme["documents"] if x["lisible_par_agent"]
                     and x["code"] not in ("declaration_etablissement",)]:
            self.deposer(troisieme, code)
        b = self.client.get(f"/dossiers/{troisieme['id']}/bilan").json()
        manquants = [l["code"] for l in b["manquants"]]
        self.assertTrue(all(c in ("echantillon", "echantillon_modele_vente", "facture_proforma") for c in manquants), manquants)

    def test_pieces_fournies_par_l_agent_depuis_la_base(self):
        # politique de la direction (rules/bibliotheque.yaml) : le fournisseur n'envoie que ses certificats
        premier = self.creer(pays="union_europeenne", fournisseur="Exemple Medical", produit="Ciment osseux Exemple (IIb)")
        for code in ("iso_13485", "piece_specifique_union_europeenne", "attestation_fabricant", "declaration_conformite",
                     "etiquetage", "notice", "catalogue"):
            self.deposer(premier, code)
        # même fabricant, même produit : l'agent fournit tout sauf les certificats
        second = self.creer(pays="union_europeenne", fournisseur="EXEMPLE MEDICAL S.p.A.", produit="Ciment osseux Exemple")
        fournis = {x["code"]: x for x in second["documents"] if x["extraction_statut"] == "terminee"}
        self.assertEqual(set(fournis), {"attestation_fabricant", "declaration_conformite", "etiquetage", "notice", "catalogue"})
        self.assertIn(f"dossier n°{premier['id']}", fournis["notice"]["origine_recu"])
        iso = self.piece(second, "iso_13485")
        self.assertIsNone(iso["extraction_statut"])  # certificats : toujours du fournisseur
        self.assertTrue(all(x["statut"] != "valide" for x in second["documents"]))  # rien n'est validé d'office
        journal = self.client.get(f"/dossiers/{second['id']}").json()["evenements"]
        self.assertTrue(any(e["action"] == "document_repris" and "fourni par l'agent" in e["detail"] for e in journal))
        # même fabricant, autre produit : seulement les pièces du fabricant
        autre_produit = self.creer(pays="union_europeenne", fournisseur="Exemple Medical", produit="Vis pédiculaire")
        self.assertEqual({x["code"] for x in autre_produit["documents"] if x["extraction_statut"] == "terminee"},
                         {"attestation_fabricant", "catalogue"})
        # autre fabricant, ou autre pays d'origine : rien
        autre = self.creer(pays="union_europeenne", fournisseur="Autre Fabricant", produit="Ciment osseux Exemple")
        self.assertFalse(any(x["extraction_statut"] == "terminee" for x in autre["documents"]))
        chine = self.creer(pays="chine", fournisseur="Exemple Medical", produit="Ciment osseux Exemple")
        self.assertFalse(any(x["extraction_statut"] == "terminee" for x in chine["documents"]))

    def test_export_du_dossier(self):
        import io
        import zipfile

        d = self.creer(pays="union_europeenne")
        self.client.post(f"/dossiers/{d['id']}/generer", json={"acteur": "Testeur"})
        self.deposer(d, "iso_13485")
        r = self.client.get(f"/dossiers/{d['id']}/export")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("BROUILLON", r.headers["content-disposition"])  # rien n'est validé
        noms = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
        self.assertEqual(noms[0], "00-Bordereau.docx")
        self.assertTrue(any(n.startswith("01-Lettre_de_demande") for n in noms), noms)
        self.assertTrue(any(n.startswith("05-Certificat_ISO_13485") and n.endswith(".pdf") for n in noms))
        bordereau = Docx(io.BytesIO(zipfile.ZipFile(io.BytesIO(r.content)).read("00-Bordereau.docx")))
        texte = "\n".join(p.text for p in bordereau.paragraphs) + "\n".join(c.text for row in bordereau.tables[0].rows for c in row.cells)
        self.assertIn("NE PAS DÉPOSER", texte)
        self.assertIn("MANQUANTE", texte)

    def test_lecture_en_echec_visible(self):
        d = self.creer()
        iso = self.piece(self.deposer(d, "iso_13485", contenu=b"illisible").json(), "iso_13485")
        self.assertEqual(iso["extraction_statut"], "erreur")
        self.assertIn("Aucun texte lisible", iso["extraction_erreur"])

    def test_depot_refuse(self):
        d = self.creer()
        self.assertEqual(self.deposer(d, "demande_signee").status_code, 409)  # pièce à rédiger
        self.assertEqual(self.deposer(d, "echantillon").status_code, 202)  # photos : reçues, rien à lire
        self.assertEqual(self.deposer(d, "echantillon_modele_vente").status_code, 409)  # objet physique, remis au dépôt
        self.assertEqual(self.deposer(d, "iso_13485", nom="virus.exe").status_code, 415)
        iso = self.piece(d, "iso_13485")
        self.client.post(f"/dossiers/{d['id']}/documents/{iso['id']}/valider",
                         json={"validateur": "Relecteur", "commentaire": "Reçu"})
        self.assertEqual(self.deposer(d, "iso_13485").status_code, 409)  # déjà validée

    def test_aucune_route_de_depot(self):
        chemins = [getattr(r, "path", "") for r in api.app.routes]
        # « document-recu » = pièce reçue DU fournisseur ; aucune route n'envoie un dossier à l'autorité
        for mot in ("depot", "deposer", "soumission", "envoi", "soumettre", "transmettre"):
            self.assertFalse(any(mot in c for c in chemins), f"route suspecte contenant '{mot}' : {chemins}")

    def test_cle_api(self):
        api.API_KEY = "secret-test"
        self.assertEqual(self.client.get("/health").status_code, 200)
        self.assertEqual(self.client.get("/dossiers").status_code, 401)
        self.assertEqual(self.client.get("/dossiers", headers={"X-API-Key": "mauvaise"}).status_code, 401)
        self.assertEqual(self.client.get("/dossiers", headers={"X-API-Key": "secret-test"}).status_code, 200)

    def test_generation_interrompue_marquee_en_erreur_au_redemarrage(self):
        d = self.creer()
        with db.SessionLocal() as s:
            doc = s.get(db.Document, self.piece(d, "demande_signee")["id"])
            doc.statut = "en_generation"
            s.commit()
        self.assertGreaterEqual(taches.reprendre_apres_redemarrage(), 1)
        d = self.client.get(f"/dossiers/{d['id']}").json()
        self.assertEqual(self.piece(d, "demande_signee")["statut"], "erreur")

    def test_dossier_inexistant(self):
        self.assertEqual(self.client.get("/dossiers/999999").status_code, 404)


if __name__ == "__main__":
    unittest.main()
