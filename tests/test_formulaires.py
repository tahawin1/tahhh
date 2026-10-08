"""
Formulaires remplis par le code (fiche signalétique, annexe II), données du
dispositif et dépôt groupé des documents du fournisseur — sans Mistral.

    python -m unittest discover -s tests
"""
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace as N

TMP = Path(tempfile.mkdtemp(prefix="conformite-formulaires-"))
os.environ.setdefault("DATABASE_URL", f"sqlite+pysqlite:///{TMP / 'test.db'}")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import pymupdf  # noqa: E402
from docx import Document as Docx  # noqa: E402

import classement  # noqa: E402
import formulaires  # noqa: E402
from rule_engine import documents_requis_maroc  # noqa: E402

PROFIL = {"raison_sociale": "SOCIETE EXEMPLE SARL", "adresse": "12 rue Exemple, Fès", "telephone": "05 00 00 00 00",
          "email": "contact@exemple.ma", "representant_legal": "M. Exemple"}


def lu(nom, valeur, verification="verifie", normalisee=None):
    return {"nom": nom, "libelle": nom, "type": "texte", "valeur": valeur, "valeur_normalisee": normalisee or valeur,
            "citation": valeur, "verification": verification}


def dossier(pays="union_europeenne", lectures=None, saisies=None):
    """Dossier factice : pièces du moteur de règles, documents « lus »."""
    documents = []
    for d in documents_requis_maroc(pays, classe="IIB"):
        champs = (lectures or {}).get(d.id)
        documents.append(N(code=d.id, numero=d.numero, nom=d.nom, nature=d.nature,
                           extraction_statut="terminee" if champs else None,
                           extraction={"champs": champs} if champs else None))
    return N(produit="Ciment 2", classe="IIB", pays_origine=pays, documents=documents, donnees_dispositif=saisies)


LECTURES_UE = {
    "piece_specifique_union_europeenne": [lu("numero", "CE 641427"), lu("numero_organisme_notifie", "2797"),
                                          lu("date_expiration", "26.05.2024", normalisee="2024-05-26"),
                                          lu("titulaire", "TECRES S.P.A.")],
    "declaration_conformite": [lu("fabricant", "TECRES S.P.A."),
                               lu("adresse_fabricant", "Via Andrea Doria 6, 37066 Sommacampagna, Italy")],
    "notice": [lu("indications", "Fixation de prothèses", verification="citation_introuvable"),
               lu("contact_fabricant", "Tel +39 045 9217311")],
    "catalogue": [lu("nom_marque", "TECRES"), lu("references", "TECRES | CEMEX RX | 1200/A\nTECRES | CEMEX FAST | 12A3000")],
}


class TestResolution(unittest.TestCase):
    def test_priorites_et_provenances(self):
        r = formulaires.resoudre(dossier(lectures=LECTURES_UE), PROFIL)
        self.assertEqual((r["designation"]["valeur"], r["designation"]["provenance"]), ("Ciment 2", "dossier"))
        self.assertEqual(r["classe"]["valeur"], "IIb")  # écrite comme sur le formulaire
        self.assertEqual((r["ce_numero"]["valeur"], r["ce_numero"]["provenance"]), ("CE 641427", "piece"))
        self.assertIn("pièce 4", r["ce_numero"]["detail"])
        self.assertEqual(r["ce_validite"]["valeur"], "26/05/2024")
        self.assertEqual(r["ce_organisme"]["valeur"], "2797")
        self.assertEqual(r["cvl_numero"]["valeur"], "NA")  # dispositif marqué CE : pas de certificat de vente libre
        self.assertEqual(r["importateur_nom"]["valeur"], "SOCIETE EXEMPLE SARL")
        self.assertEqual(r["fabricant_nom"]["valeur"], "TECRES S.P.A.")
        self.assertEqual(r["site_fabrication"]["valeur"], "TECRES S.P.A.")  # repris d'une autre donnée
        self.assertEqual(r["fabrication"]["valeur"], "Importé")
        self.assertTrue(r["objet"]["a_verifier"])  # valeur habituelle : à confirmer
        self.assertTrue(r["indications"]["a_verifier"])  # citation introuvable dans la notice
        self.assertFalse(r["ce_numero"]["a_verifier"])

    def test_rien_n_est_invente(self):
        r = formulaires.resoudre(dossier(), {})
        for ident in ("indications", "nom_marque", "fabricant_nom", "contient_latex", "importateur_nom", "ce_numero"):
            self.assertIsNone(r[ident]["valeur"], ident)
            self.assertEqual(r[ident]["provenance"], "manquant")

    def test_saisie_prioritaire(self):
        r = formulaires.resoudre(dossier(lectures=LECTURES_UE, saisies={"fabricant_nom": "TECRES SpA"}), PROFIL)
        self.assertEqual((r["fabricant_nom"]["valeur"], r["fabricant_nom"]["provenance"]), ("TECRES SpA", "saisie"))

    def test_sterilisation_et_validite_relevees_par_le_code(self):
        d = dossier("inde", lectures={"etiquetage": [lu("sterilisation", "STERILE")]})
        etiquette = next(x for x in d.documents if x.code == "etiquetage")
        etiquette.texte_recu = "SURGICAL GLOVES  STERILE EO  Shelf Life : 5 Years  Sterilized using Ethylene Oxide"
        r = formulaires.resoudre(d, PROFIL)
        self.assertEqual(r["sterilisation"]["valeur"], "Oxyde d'éthylène")
        self.assertEqual(r["duree_validite"]["valeur"], "5 ans")
        self.assertTrue(r["sterilisation"]["a_verifier"])
        etiquette.texte_recu = "Gloves. Use before expiry date."  # rien de concluant : « STERILE » seul rejeté
        r = formulaires.resoudre(d, PROFIL)
        self.assertIsNone(r["sterilisation"]["valeur"])
        self.assertIsNone(r["duree_validite"]["valeur"])
        notice = next(x for x in d.documents if x.code == "notice")
        notice.texte_recu = "STERILE R. Gamma irradiated. 3 years from the date of manufacture."
        r = formulaires.resoudre(d, PROFIL)
        self.assertIn("Irradiation", r["sterilisation"]["valeur"])
        self.assertEqual(r["duree_validite"]["valeur"], "3 ans")

    def test_etiquette_chinoise_duree_presentation_references_contact(self):
        d = dossier("chine", lectures={
            "etiquetage": [lu("duree_validite", "Expiration date: 2023-11-08"),
                           lu("references", "MICROCURE | WZDSS-A-35W [LOT 2409015 Qty:10 |\n"
                                            "MICROCURE | WZDSS-A-35W [LOT 2409016 Qty:10 |\nMICROCURE | REF: WZDSS-A-15N")],
            "declaration_conformite": [lu("contact_fabricant", "Tel: 86-512-62916116")]})
        etiquette = next(x for x in d.documents if x.code == "etiquetage")
        etiquette.texte_recu = ("Disposable Skin Stapler STERILE EO Single use. Qty: 10 "
                                "MFG 2024-09-05 EXP 2027-09-04")
        next(x for x in d.documents if x.code == "declaration_conformite").texte_recu = "E-mail: info@microcure.com.cn"
        r = formulaires.resoudre(d, PROFIL)
        self.assertEqual(r["duree_validite"]["valeur"], "3 ans")  # fabrication -> péremption, pas une date seule
        self.assertEqual(r["presentation"]["valeur"], "Emballage unitaire stérile/Boite de 10 unités stériles")
        lignes = r["references"]["valeur"].splitlines()  # lot et quantité retirés, doublons fusionnés
        self.assertEqual([l.split("|")[-1].strip() for l in lignes], ["WZDSS-A-35W", "WZDSS-A-15N"])
        self.assertIn("info@microcure.com.cn", r["fabricant_contact"]["valeur"])
        self.assertEqual(formulaires.duree_en_francais("2024.09.05 ag 2027.09.04"), "3 ans")
        self.assertEqual(formulaires.duree_en_francais("Shelf life 18 months"), "18 mois")

    def test_methode_la_plus_citee_dans_tout_le_dossier(self):
        d = dossier("inde", lectures={"etiquetage": [lu("duree_validite", "Pas spécifié")]})
        next(x for x in d.documents if x.code == "etiquetage").texte_recu = "Can be sterilized by steam."
        next(x for x in d.documents if x.code == "notice").texte_recu = "STERILE EO. Sterilized using ethylene oxide."
        r = formulaires.resoudre(d, PROFIL)
        self.assertEqual(r["sterilisation"]["valeur"], "Oxyde d'éthylène")
        self.assertIsNone(r["duree_validite"]["valeur"])  # « Pas spécifié » n'est pas une valeur

    def test_un_nom_commercial_plusieurs_references(self):
        d = dossier("chine", lectures={"catalogue": [lu("nom_marque", "MICROCURE"), lu(
            "references", "Disposable Skin Stapler | WZDSS-A-35W\nDisposable Skin Stapler | WZDSS-A-15N")]})
        r = formulaires.resoudre(d, PROFIL)
        self.assertEqual(r["nom_marque"]["valeur"], "MICROCURE / Disposable Skin Stapler")

    def test_usages_des_dossiers_acceptes(self):
        d = dossier("inde", lectures={
            "catalogue": [lu("nom_marque", "iLife Medical Devices"),
                          lu("references", "ILIFE MEDICAL DEVICES | OneFlon Infusion Set | 903.01.02.00")],
            "declaration_conformite": [lu("fabricant", "iLife Medical Devices Pvt. Ltd."),
                                       lu("cl_declaration_de_conformite_0", "EMDN A03040199")]})
        d.produit, d.preuve = "Set de perfusion stérile", "ce"
        next(x for x in d.documents if x.code == "etiquetage").texte_recu = "STERILE EO  CE 0123  Single use"
        r = formulaires.resoudre(d, PROFIL)
        self.assertEqual(r["categorie"]["valeur"], "Injection-Prélèvement")
        self.assertTrue(r["categorie"]["a_verifier"])
        self.assertEqual(r["code_classification"]["valeur"], "CE 0123 / A03040199")
        self.assertEqual(r["nom_marque"]["valeur"], "ILIFE / OneFlon Infusion Set")
        d.produit = "Set de transfusion sanguine"
        r = formulaires.resoudre(d, PROFIL)
        self.assertEqual((r["categorie"]["valeur"], r["domaine_therapeutique"]["valeur"]),
                         ("Sang et circulation sanguine", "Transfusion sanguine"))
        d.produit = "Produit inconnu"
        self.assertIsNone(formulaires.resoudre(d, PROFIL)["categorie"]["valeur"])  # rien d'inventé

    def test_textes_en_francais(self):
        notice = [lu("indications", "Used to deliver blood"), lu("composition", "Composizione dei cementi ossei CEMEX:"),
                  lu("conservation", "Store in a cool and dry place. Protect from direct sunlight.")]
        notice[0]["valeur_fr"] = "Utilisé pour administrer du sang"
        r = formulaires.resoudre(dossier("inde", lectures={"notice": notice}), PROFIL)
        self.assertEqual(r["indications"]["valeur"], "Utilisé pour administrer du sang")
        self.assertTrue(r["indications"]["a_verifier"])
        self.assertIn("Used to deliver blood", r["indications"]["detail"])  # texte d'origine affiché
        self.assertEqual(r["conservation"]["valeur"], "Conservé à l’abri de la chaleur et de l’humidité")
        self.assertIn("Store in a cool", r["conservation"]["detail"])
        self.assertEqual(r["composition"]["valeur"], "Ciment 2")  # un intitulé n'est pas une composition
        self.assertTrue(r["composition"]["a_verifier"])

    def test_conservation_relevee_dans_le_texte(self):
        d = dossier("inde")
        next(x for x in d.documents if x.code == "notice").texte_recu = (
            "Sterile set. Heat sealed pouch. Storage: keep in a dry place, away from sunlight.")
        self.assertEqual(formulaires.resoudre(d, PROFIL)["conservation"]["valeur"],
                         "Conservé à l’abri de la chaleur et de l’humidité")
        next(x for x in d.documents if x.code == "notice").texte_recu = "Heat sealed pouch. Do not reuse."
        self.assertIsNone(formulaires.resoudre(d, PROFIL)["conservation"]["valeur"])

    def test_origine_chine_certificat_ce_sans_objet(self):
        r = formulaires.resoudre(dossier("chine"), PROFIL)
        self.assertEqual(r["ce_numero"]["valeur"], "NA")
        self.assertIsNone(r["cvl_numero"]["valeur"])

    def test_fabricant_chinois_marque_ce(self):
        # même pays, preuve CE : les cases du certificat CE se remplissent, celles du CVL sont sans objet
        d = dossier("chine")
        d.preuve = "ce"
        r = formulaires.resoudre(d, PROFIL)
        self.assertNotEqual(r["ce_numero"]["valeur"], "NA")
        self.assertEqual(r["cvl_numero"]["valeur"], "NA")
        self.assertIn("marqué CE", r["cvl_numero"]["detail"])

    def test_reference_2023_607_n_est_pas_le_numero_ce(self):
        # essai réel : l'agent a lu « EU2023-607/… » (référence de la lettre) comme n° de certificat
        lettre = ("Notified Body Confirmation Letter Reference: EU2023-607/123456 Regulation (EU) 2023/607\n"
                  "Exemple Cement manual\nB-UDI 0000000001\nClass IIb implantable non-\nWET\nN/A\n"
                  "Certificate #1 CE 641427; NB (2797)\nExemple Prep Kit\nB-UDI 0000000002\n"
                  "Class IIb excluding Class\nIIb implantable non-WET\nN/A\nCertificate #1 CE 641427; NB (2797)\n")
        d = dossier(lectures={"piece_specifique_union_europeenne": [lu("numero", "EU2023-607/123456")],
                              "catalogue": [lu("nom_marque", "Exemple"), lu("references", "Exemple Cement manual")]})
        for p in d.documents:
            p.texte_recu = lettre if p.code.startswith("piece_specifique") else ""
        r = formulaires.resoudre(d, PROFIL)
        self.assertEqual(r["ce_numero"]["valeur"], "CE 641427")  # relevé dans le texte, à vérifier
        self.assertTrue(r["ce_numero"]["a_verifier"])
        # fin de transition 2023/607 lue sur la ligne du produit : IIb implantable -> 31/12/2027
        self.assertIn("31/12/2027", r["ce_validite"]["valeur"])
        self.assertNotIn("2028", r["ce_validite"]["valeur"])
        self.assertIn("Exemple Cement manual", r["ce_validite"]["detail"])
        # produit ambigu (gamme aux deux dates) : les deux dates, à préciser
        d.documents[[p.code for p in d.documents].index("catalogue")].extraction = {
            "champs": [lu("nom_marque", "Exemple"), lu("references", "Exemple")]}
        r = formulaires.resoudre(d, PROFIL)
        self.assertIn("à préciser", r["ce_validite"]["valeur"])

    def test_plusieurs_produits_en_annexe(self):
        d = dossier(lectures={"catalogue": [lu("nom_marque", "Tecres"),
                                            lu("references", "Cemex RX, Cemex Fast, Cemex Isoplastic")]})
        r = formulaires.resoudre(d, PROFIL)
        self.assertEqual(r["nom_marque"]["valeur"], "TECRES / Voir annexe")  # comme le dossier accepté
        self.assertEqual(formulaires._lignes_references(r["references"]["valeur"]),
                         [["TECRES", "Cemex RX", ""], ["TECRES", "Cemex Fast", ""], ["TECRES", "Cemex Isoplastic", ""]])
        # les étiquettes priment sur le catalogue : produits déposés, avec leur REF
        d = dossier(lectures={"catalogue": [lu("nom_marque", "Tecres"), lu("references", "Cemex RX, Vancogenx")],
                              "etiquetage": [lu("references", "CEMEX RX | 1200/A\nCEMEX FAST | 12A3000")]})
        r = formulaires.resoudre(d, PROFIL)
        self.assertEqual(formulaires._lignes_references(r["references"]["valeur"]),
                         [["TECRES", "CEMEX RX", "1200/A"], ["TECRES", "CEMEX FAST", "12A3000"]])
        self.assertEqual(r["nom_marque"]["valeur"], "TECRES / Voir annexe")

    def test_toutes_les_lignes_des_formulaires_sont_definies(self):
        regles = formulaires.charger()
        for code, f in regles["formulaires"].items():
            for s in f["sections"]:
                for ligne in s["lignes"]:
                    ident = next(iter(ligne)) if isinstance(ligne, dict) else ligne
                    self.assertIn(ident, regles["donnees"], f"{code} : {ident}")
        self.assertEqual(formulaires.formulaires(), {"fiche_signaletique", "certificat_enregistrement_annexe2"})


class TestDocx(unittest.TestCase):
    def tearDown(self):
        shutil.rmtree(TMP / "sortie", ignore_errors=True)

    def texte(self, chemin):
        d = Docx(chemin)
        cellules = [c.text for t in d.tables for row in t.rows for c in row.cells]
        return "\n".join([p.text for p in d.paragraphs] + cellules), d

    def test_fiche_signaletique(self):
        res = formulaires.generer_formulaire(dossier(lectures=LECTURES_UE), "fiche_signaletique", "Fiche signalétique",
                                            TMP / "sortie", PROFIL)
        texte, d = self.texte(res["fichier"])
        for attendu in ("FICHE SIGNALÉTIQUE", "1- OBJET DE LA DEMANDE", "☒ Nouvel enregistrement", "☐ Mise à jour",
                        "☒ Importé", "CE 641427", "2797", "SOCIETE EXEMPLE SARL", "TECRES S.P.A.",
                        "CEMEX FAST", "12A3000", "[À COMPLÉTER]", "M. Exemple", "EN ATTENTE DE VALIDATION HUMAINE"):
            self.assertIn(attendu, texte)
        self.assertGreater(len(d.tables), 7)  # une table par section + l'annexe
        self.assertGreater(res["bilan"]["remplies"], 20)
        self.assertIn("rempli par le code", res["sources"][0]["texte_source"])

    def test_annexe2_partie_ministere_jamais_remplie(self):
        res = formulaires.generer_formulaire(dossier(lectures=LECTURES_UE), "certificat_enregistrement_annexe2",
                                            "Annexe II", TMP / "sortie", PROFIL)
        texte, _ = self.texte(res["fichier"])
        self.assertIn("Code CE (pour les dispositifs médicaux marqués CE)", texte)
        self.assertIn("ne pas remplir", texte)
        self.assertNotIn("Suite à l'avis de la Commission", texte)


def pdf(texte):
    d = pymupdf.open()
    d.new_page().insert_text((72, 72), texte)
    return d.tobytes()


class TestClassement(unittest.TestCase):
    PIECES = [N(code=d.id, numero=d.numero, nom=d.nom) for d in documents_requis_maroc("union_europeenne", "IIB")
              if d.nature == "a_fournir"]

    def numero(self, nom):
        p, _ = classement.classer(nom, self.PIECES)
        return p.numero if p else None

    def test_noms_du_dossier_accepte(self):
        attendus = {
            "3-3-Authorization Letter 2026.pdf": 3, "Confirmation Letter TECRES (ID 658919).pdf": 4,
            "EC_Certificate_FQA_CE641427_26.02.2024.pdf": 4, "ISO13485.pdf": 5, "Certificat CE ISO 13485.pdf": 5,
            "DoC for Morocco_Cemex_DRAFT (1).pdf": 6, "1200A_Cemex RX_Label.pdf": 7, "IFU PAGE 27.pdf": 8,
            "CATALOGUE.pdf": 14, "facture.pdf": None,
        }
        for nom, numero in attendus.items():
            self.assertEqual(self.numero(nom), numero, nom)

    def test_plusieurs_fichiers_reunis_en_un_pdf(self):
        nom, extension, contenu = classement.reunir([("b.pdf", ".pdf", pdf("lettre")), ("a.pdf", ".pdf", pdf("certificat"))])
        self.assertEqual(extension, ".pdf")
        d = pymupdf.open(stream=contenu, filetype="pdf")
        self.assertEqual([p.get_text().strip() for p in d], ["certificat", "lettre"])  # dans l'ordre des noms
        self.assertEqual(nom, "a.pdf + b.pdf.pdf")


if __name__ == "__main__":
    unittest.main()
