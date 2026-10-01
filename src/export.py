"""
Préparation du dépôt : le dossier en un fichier ZIP, pièces dans l'ordre du
dossier déposé (« 01-… », « 02-… », …) et bordereau des pièces.

Le dépôt lui-même reste manuel (AMMPS, mercredi ou jeudi). Tant que toutes les
pièces ne sont pas validées par une personne, le ZIP est marqué BROUILLON et
le bordereau porte « ne pas déposer ».
"""
from __future__ import annotations

import re
import unicodedata
import zipfile
from pathlib import Path

from rule_engine import prochain_creneau_depot


def _slug(texte: str, longueur: int = 60) -> str:
    t = unicodedata.normalize("NFKD", texte)
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^A-Za-z0-9]+", "_", t).strip("_")[:longueur] or "piece"


def _fichier_de(piece) -> Path | None:
    chemin = piece.fichier if piece.nature == "a_rediger" else piece.fichier_recu
    return Path(chemin) if chemin and Path(chemin).exists() else None


def construire(dossier) -> tuple[Path, bool]:
    """Écrit le ZIP du dossier dans son dossier de sortie ; retourne (chemin, complet)."""
    from docx import Document

    pieces = sorted(dossier.documents, key=lambda p: (p.numero is None, p.numero or 0, p.ordre))
    complet = bool(pieces) and all(p.statut == "valide" for p in pieces)
    lignes, contenu = [], []
    rang_libre = 90  # pièces sans numéro : après les pièces numérotées
    for p in pieces:
        if p.numero is not None:
            prefixe = f"{p.numero:02d}"
        else:
            rang_libre += 1
            prefixe = f"{rang_libre}"
        fichier = _fichier_de(p)
        nom = f"{prefixe}-{_slug(p.nom)}{fichier.suffix.lower() if fichier else ''}"
        if fichier:
            contenu.append((fichier, nom))
        if p.statut == "valide":
            etat = f"Validée par {p.valide_par} le {p.valide_le:%d/%m/%Y}" if p.valide_le else "Validée"
        elif fichier:
            etat = "À valider"
        else:
            etat = "MANQUANTE"
        lignes.append((p.numero, p.nom, nom if fichier else "—", etat))

    sortie = Path(dossier.dossier_sortie)
    sortie.mkdir(parents=True, exist_ok=True)
    bordereau = sortie / "00-Bordereau.docx"
    doc = Document()
    doc.add_heading("Bordereau des pièces — demande d'enregistrement d'un dispositif médical", level=1)
    doc.add_paragraph(f"Dispositif : {dossier.produit} — classe {dossier.classe or 'non précisée'} — "
                      f"origine : {dossier.pays_origine} — dossier n° {dossier.id}")
    if complet:
        doc.add_paragraph().add_run(
            f"DOSSIER COMPLET — toutes les pièces validées. Dépôt manuel à l'AMMPS (mercredi ou jeudi) ; "
            f"prochain créneau : {prochain_creneau_depot():%d/%m/%Y}.").bold = True
    else:
        doc.add_paragraph().add_run(
            "BROUILLON — dossier non entièrement validé : NE PAS DÉPOSER.").bold = True
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cellule, titre in zip(table.rows[0].cells, ("N°", "Pièce", "Fichier", "État")):
        cellule.paragraphs[0].add_run(titre).bold = True
    for numero, nom, fichier, etat in lignes:
        for cellule, texte in zip(table.add_row().cells, (str(numero) if numero is not None else "—", nom, fichier, etat)):
            cellule.text = texte
    doc.save(bordereau)

    archive = sortie / f"dossier{dossier.id}_{_slug(dossier.produit, 40)}{'' if complet else '_BROUILLON'}.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(bordereau, bordereau.name)
        for fichier, nom in contenu:
            z.write(fichier, nom)
    return archive, complet
