"""
Formulaires officiels remplis par le code — fiche signalétique (pièce 2) et
certificat d'enregistrement de l'annexe II (pièce 16).

Mistral ne respecte pas la mise en page d'un formulaire : ces pièces ne lui
sont pas confiées. Chaque case est remplie à partir d'une source déclarée dans
rules/formulaires_maroc.yaml, dans cet ordre de priorité :

    saisie de l'utilisateur > document reçu lu par l'agent > profil de
    l'entreprise > données du dossier > autre donnée > valeur habituelle

et chaque valeur garde sa provenance, affichée à la validation. Aucune valeur
n'est inventée : sans source, la case reste [À COMPLÉTER].
"""
from __future__ import annotations

import datetime
import re
from pathlib import Path

import yaml

from rule_engine import RULES_DIR

FICHIER = RULES_DIR / "formulaires_maroc.yaml"
A_COMPLETER = "[À COMPLÉTER]"
CLASSES = {"I": "I", "IS": "Is", "IM": "Im", "IR": "Ir", "IIA": "IIa", "IIB": "IIb", "III": "III"}
MENTION_VALIDATION = (
    "PROJET GÉNÉRÉ AUTOMATIQUEMENT — EN ATTENTE DE VALIDATION HUMAINE. "
    "Ne pas déposer avant relecture et validation explicite."
)


SANS_VALEUR = re.compile(r"^\W*(?:pas|non|not)?\s*(?:sp[ée]cifi[ée]e?|specified|indiqu[ée]e?|mentionn[ée]e?|stated|"
                         r"available|disponible|applicable)\W*$|^\W*(?:n/?a|none|null|inconnu|unknown|aucune?)\W*$", re.I)


def charger() -> dict:
    with open(FICHIER, encoding="utf-8") as f:
        return yaml.safe_load(f)


def formulaires() -> set[str]:
    """Codes des pièces remplies par le code plutôt que rédigées par Mistral."""
    return set(charger()["formulaires"])


def _lectures(documents) -> dict[str, dict]:
    """{code de pièce : {champ : champ lu}} pour les documents reçus déjà lus."""
    lectures = {}
    for d in documents:
        if d.extraction_statut == "terminee" and d.extraction:
            lectures[d.code] = {c["nom"]: {**c, "piece": d.numero, "piece_nom": d.nom}
                                for c in d.extraction.get("champs", [])}
    return lectures


def _lu(lectures: dict, code: str, champ: str) -> dict | None:
    candidats = [c for c in lectures if c == code or (code == "piece_specifique" and c.startswith("piece_specifique"))]
    for c in candidats:
        lu = lectures[c].get(champ)
        if lu and lu.get("valeur"):
            return lu
    return None


DATE = r"(?:(\d{4})[-./](\d{1,2})[-./](\d{1,2})|(\d{1,2})[-./](\d{1,2})[-./](\d{4}))"


def _dates(texte: str) -> list[datetime.date]:
    trouvees = []
    for m in re.finditer(DATE, texte or ""):
        a, mo, j = (m.group(1), m.group(2), m.group(3)) if m.group(1) else (m.group(6), m.group(5), m.group(4))
        try:
            trouvees.append(datetime.date(int(a), int(mo), int(j)))
        except ValueError:
            continue
    return trouvees


def duree_entre(debut: datetime.date, fin: datetime.date) -> str | None:
    """2024-09-05 -> 2027-09-04 : « 3 ans » (fabrication -> péremption)."""
    mois = round((fin - debut).days / 30.44)
    if mois <= 0:
        return None
    return f"{mois // 12} ans" if mois % 12 == 0 else f"{mois} mois"


def duree_en_francais(valeur: str) -> str:
    """« 5 years » -> « 5 ans » ; « 2024.09.05 à 2027.09.04 » -> « 3 ans »."""
    if (m := re.search(r"\b(\d{1,2})\s*(years?|yrs?|ans?|months?|mois)\b", valeur, re.I)):
        return f"{int(m.group(1))} {'mois' if m.group(2).lower().startswith('m') else 'ans'}"
    dates = _dates(valeur)
    if len(dates) >= 2 and (duree := duree_entre(dates[0], dates[1])):
        return duree
    return valeur


def _formater(definition: dict, valeur: str, normalisee: str | None = None) -> str:
    if definition.get("format") == "duree":
        return duree_en_francais(valeur)
    if definition.get("format") == "classe":
        return CLASSES.get(valeur.upper(), valeur)
    if definition.get("format") == "date" and normalisee:
        try:
            return datetime.date.fromisoformat(normalisee).strftime("%d/%m/%Y")
        except ValueError:
            pass
    return valeur


def resoudre(dossier, profil: dict, saisies: dict | None = None) -> dict[str, dict]:
    """Valeur et provenance de chaque donnée du dispositif pour ce dossier.
    `dossier` : objet avec produit, classe, pays_origine, documents."""
    definitions = charger()["donnees"]
    saisies = (saisies if saisies is not None else getattr(dossier, "donnees_dispositif", None)) or {}
    lectures = _lectures(dossier.documents)
    contexte = {"produit": dossier.produit, "classe": dossier.classe or "", "pays_origine": dossier.pays_origine}
    # fabricant non européen marqué CE : les cases du certificat CE s'appliquent, celles du CVL non
    from rule_engine import pays_des_regles

    pays_regles = pays_des_regles(dossier.pays_origine, getattr(dossier, "preuve", None))
    resultat: dict[str, dict] = {}

    def une(ident: str, pile: tuple = ()) -> dict:
        if ident in resultat:
            return resultat[ident]
        d = definitions[ident]
        r = {"id": ident, "libelle": d["libelle"], "type": d.get("type", "texte"), "options": d.get("options"),
             "valeur": None, "provenance": "manquant", "detail": None, "a_verifier": False}
        if (saisie := (saisies.get(ident) or "").strip()):
            r.update(valeur=saisie, provenance="saisie", detail="Saisi dans le tableau de bord")
        elif pays_regles in (d.get("na_si_origine") or []):
            r.update(valeur="NA", provenance="regle", detail=f"Sans objet pour l'origine {pays_regles}"
                     + (" (fabricant marqué CE)" if pays_regles != dossier.pays_origine else ""))
        else:
            for source in d.get("sources") or []:
                genre, _, reste = source.partition(":")
                if genre == "dossier" and contexte.get(reste):
                    r.update(valeur=_formater(d, contexte[reste]), provenance="dossier", detail="Données du dossier")
                elif genre == "profil" and profil.get(reste):
                    r.update(valeur=profil[reste], provenance="profil", detail="Profil de l'entreprise")
                elif genre == "piece":
                    code, _, champ = reste.partition(":")
                    lu = _lu(lectures, code, champ)
                    if lu and SANS_VALEUR.match(lu["valeur"]):
                        lu = None  # « Pas spécifié », « N/A » : Mistral n'a rien trouvé
                    if lu and any(re.search(motif, lu["valeur"]) for motif in d.get("rejeter") or []):
                        lu = None  # ex. référence de la lettre 2023/607 lue à la place du n° de certificat CE
                    if lu:
                        verifie = lu.get("verification") == "verifie"
                        traduit = bool(lu.get("valeur_fr"))
                        r.update(valeur=lu["valeur_fr"] if traduit else _formater(d, lu["valeur"], lu.get("valeur_normalisee")),
                                 provenance="piece", a_verifier=not verifie or traduit,
                                 detail=f"Lu dans la pièce {lu['piece']} ({lu['piece_nom']})"
                                        + ("" if verifie else " — non vérifié dans le texte")
                                        + (f" — traduit en français par Mistral, texte d'origine : « {lu['valeur'][:200]} »"
                                           if traduit else ""))
                elif genre == "motif":
                    code, _, motif = reste.partition(":")
                    if (trouve := _motif(dossier.documents, code, motif)):
                        valeur, piece = trouve
                        r.update(valeur=valeur, provenance="piece", a_verifier=True,
                                 detail=f"Relevé dans le texte de la pièce {piece.numero} ({piece.nom})")
                elif genre == "references_etiquettes":
                    if (trouve := _references_etiquettes(dossier.documents)):
                        valeur, piece = trouve
                        r.update(valeur=valeur, provenance="piece", a_verifier=True,
                                 detail=f"Relevé par le code sur les étiquettes (pièce {piece.numero})")
                elif genre == "par_produit":
                    regle = next((x for x in d.get("par_produit") or []
                                  if re.search(x["si"], dossier.produit or "", re.I)), None)
                    if regle:
                        r.update(valeur=regle["valeur"], provenance="defaut", a_verifier=True,
                                 detail=f"Valeur des dossiers acceptés pour ce type de produit (« {dossier.produit} ») "
                                        "— à confirmer")
                elif genre in RELEVES:
                    if (trouve := RELEVES[genre](dossier.documents)):
                        valeur, piece = trouve
                        r.update(valeur=valeur, provenance="piece", a_verifier=True,
                                 detail=f"Relevé par le code dans le texte de la pièce {piece.numero} ({piece.nom})")
                elif genre == "transition_2023_607":
                    if (trouve := _transition_2023_607(dossier, une)):
                        valeur, detail = trouve
                        r.update(valeur=valeur, provenance="piece", a_verifier=True, detail=detail)
                elif genre == "donnee" and reste not in pile:
                    autre = une(reste, pile + (ident,))
                    if autre["valeur"]:
                        r.update(valeur=autre["valeur"], provenance="donnee", a_verifier=autre["a_verifier"],
                                 detail=f"Repris de « {autre['libelle']} »")
                elif genre == "defaut":
                    r.update(valeur=reste, provenance="defaut", a_verifier=True,
                             detail="Valeur habituelle (dossier accepté) — à confirmer")
                if r["valeur"] and d.get("garder"):  # ne garder que le code (« EMDN code P099001- » -> P099001)
                    m = re.search(d["garder"], r["valeur"])
                    r["valeur"] = m.group(0) if m else None
                if r["valeur"]:
                    break
        if d.get("a_confirmer") and r["valeur"] and r["provenance"] != "saisie":
            r["a_verifier"] = True
        if (habituelle := d.get("habituelle")) and r["provenance"] in ("piece", "manquant"):
            _formule_habituelle(r, habituelle, dossier.documents)
        resultat[ident] = r
        return r

    for ident in definitions:
        une(ident)
    code = resultat.get("code_classification")
    if code and code["provenance"] != "saisie" and definitions["code_classification"].get("organisme_notifie") \
            and pays_regles == "union_europeenne" and (organisme := _organisme_notifie(dossier.documents, une("ce_organisme"))):
        code.update(valeur=f"CE {organisme}" + (f" / {code['valeur']}" if code["valeur"] else ""), a_verifier=True,
                    provenance=code["provenance"] if code["valeur"] else "piece",
                    detail=f"N° de l'organisme notifié en tête, comme les dossiers acceptés"
                           + (f" ; {code['detail']}" if code["detail"] else ""))
    contact = resultat.get("fabricant_contact")
    if contact and contact["valeur"] and contact["provenance"] == "piece" and "@" not in contact["valeur"] \
            and (trouve := _courriel(dossier.documents)):
        contact["valeur"] += f"\nAdresse électronique : {trouve[0]}"
        contact["a_verifier"] = True
        contact["detail"] = (contact["detail"] or "") + f" — adresse électronique relevée dans la pièce {trouve[1].numero}"
    _memoire(resultat, dossier)
    _plusieurs_produits(resultat)
    return resultat


def _formule_habituelle(r: dict, habituelle: dict, documents) -> None:
    """Formule employée par la majorité des dossiers acceptés (ex. conservation :
    « Conservé à l’abri de la chaleur et de l’humidité ») quand le texte du
    fournisseur dit la même chose (« Store in a cool, dry place… ») — à confirmer,
    le texte du fournisseur restant affiché."""
    motif = re.compile(habituelle["si"], re.I)
    source = r["valeur"]
    if not source and habituelle.get("phrase"):  # rien de lu : une phrase du texte reçu qui parle de stockage
        phrase = re.compile(habituelle["phrase"], re.I)
        source = next((p for code in habituelle.get("pieces", ()) for _, texte in _textes(documents, code)
                       for p in re.split(r"(?<=[.;])\s+|\n", texte) if phrase.search(p) and motif.search(p)), None)
    if source and motif.search(source):
        r.update(valeur=habituelle["valeur"], provenance="piece" if r["valeur"] else "regle", a_verifier=True,
                 detail=f"Formule habituelle des dossiers acceptés ; le fournisseur écrit : « {source.strip()[:200]} »")


def _memoire(resultat: dict, dossier) -> None:
    """Cases encore vides (ou à la valeur habituelle) : reprises du dossier
    accepté du MÊME fabricant (mémoire apprise par Mistral, voir memoire.py),
    toujours à confirmer. Une saisie ou une lecture du document reçu prime."""
    import memoire

    fabricant = resultat.get("fabricant_nom", {}).get("valeur") or getattr(dossier, "fournisseur", None)
    accepte = memoire.pour(fabricant, dossier.pays_origine, dossier.produit)
    if accepte:
        for ident, appris in accepte["valeurs"].items():
            r = resultat.get(ident)
            if r is None or r["provenance"] not in ("manquant", "defaut"):
                continue
            r.update(valeur=appris["valeur"], provenance="memoire", a_verifier=True,
                     detail=f"Repris du dossier accepté « {accepte['produit']} » (pièce {appris['piece']}) — à confirmer")
    # catégorie et domaine thérapeutique : propres au produit, repris du produit accepté le plus proche
    if any(resultat.get(i, {}).get("provenance") == "manquant" for i in memoire.CHAMPS_DU_PRODUIT) \
            and (proche := memoire.pour_produit(dossier.produit)):
        for ident in memoire.CHAMPS_DU_PRODUIT:
            r, appris = resultat.get(ident), proche["valeurs"].get(ident)
            if r is not None and appris and r["provenance"] == "manquant":
                r.update(valeur=appris["valeur"], provenance="memoire", a_verifier=True,
                         detail=f"Repris du produit accepté le plus proche « {proche['produit']} » "
                                f"({proche.get('fabricant') or 'autre fabricant'}) — à confirmer")


def _textes(documents, code: str):
    for d in documents:
        if (d.code == code or (code == "piece_specifique" and d.code.startswith("piece_specifique"))) \
                and getattr(d, "texte_recu", None):
            yield d, d.texte_recu


MENTIONS_ETIQUETTE = {"FOR", "INTERNAL", "USE", "ONLY", "LOT", "REF", "STERILE", "CE", "EXTERNAL"}
REF_ETIQUETTE = re.compile(r"([A-Z][A-Za-z0-9®’'\-+ ]{2,40}?)\s*\(\s*Ref(?:\.|erence)?\s*(?:Code|No\.?)?\s*[:.]?\s*"
                           r"([A-Z0-9][A-Z0-9/\-.]{2,14})\s*\)")


def _references_etiquettes(documents) -> tuple[str, object] | None:
    """« CEMEX RX (Ref. Code 1200/A) » sur chaque étiquette : « NOM | RÉF »,
    une ligne par produit, dans l'ordre, sans doublon."""
    for piece, texte in _textes(documents, "etiquetage"):
        vus = {}
        for m in REF_ETIQUETTE.finditer(re.sub(r"\s+", " ", texte)):
            mots = m.group(1).split()[-4:]  # le nom commercial : les derniers mots avant « (Ref. … »
            while mots and mots[0].upper() in MENTIONS_ETIQUETTE:
                mots.pop(0)
            if mots:
                vus.setdefault(m.group(2).rstrip(".-"), " ".join(mots))
        if vus:
            return "\n".join(f"{nom} | {ref}" for ref, nom in vus.items()), piece
    return None


def _motif(documents, code: str, motif: str) -> tuple[str, object] | None:
    """Valeur la plus fréquente correspondant à `motif` dans le texte reçu
    (ex. « CE 641427 », répété dans la lettre 2023/607) — relevé par le code."""
    from collections import Counter
    for piece, texte in _textes(documents, code):
        trouves = Counter(re.sub(r"\s+", " ", m.group(0)).strip() for m in re.finditer(motif, texte))
        if trouves:
            return trouves.most_common(1)[0][0], piece
    return None


# Méthode de stérilisation : symboles ISO 15223 (« STERILE EO », « STERILE R »)
# et mentions du texte. La plus citée l'emporte ; « STERILE » seul ne dit rien.
METHODES_STERILISATION = (
    ("Oxyde d'éthylène", re.compile(r"ethylene[ -]?oxide|oxyde d.?[ée]thyl[èe]ne|\bSTERILE\s*[|/-]?\s*E[O0]\b|"
                                    r"\bEt[O0]\b|\bEO\s+steril", re.I)),
    ("Irradiation (rayonnement gamma)", re.compile(r"gamma[ -]?(?:ray|irradiat|radiat|steril|rayon)|rayonnement gamma|"
                                                    r"rayons gamma", re.I)),
    ("Irradiation", re.compile(r"\bSTERILE\s*[|/-]?\s*R\b|sterili[sz]ed (?:by|using|with) (?:irradiation|radiation)|"
                               r"st[ée]rilis[ée] par irradiation", re.I)),
    ("Vapeur d'eau (chaleur humide)", re.compile(r"sterili[sz]ed (?:by|using|with) (?:steam|moist heat)|"
                                                 r"st[ée]rilis[ée] (?:à|par) la vapeur", re.I)),
    ("Traitement aseptique", re.compile(r"\bSTERILE\s*[|/-]?\s*A\b|aseptic(?:ally)? process", re.I)),
)
PIECES_ETIQUETAGE = ("etiquetage", "notice", "catalogue", "declaration_conformite")


def _sterilisation(documents) -> tuple[str, object] | None:
    """Méthode la plus citée dans l'ensemble des étiquettes, notice, catalogue et
    déclaration (une mention isolée dans une pièce ne l'emporte pas)."""
    comptes, pieces = {}, {}
    for code in PIECES_ETIQUETAGE:
        for piece, texte in _textes(documents, code):
            for nom, motif in METHODES_STERILISATION:
                if (n := len(motif.findall(texte))):
                    comptes[nom] = comptes.get(nom, 0) + n
                    pieces.setdefault(nom, piece)
    if not comptes:
        return None
    nom = max(comptes, key=comptes.get)
    return nom, pieces[nom]


DUREE = re.compile(
    r"(?:shelf[- ]?life|validity|validit[ée]|dur[ée]e de (?:validit[ée]|conservation)|expir\w*|p[ée]remption)"
    r"[^.\n]{0,40}?\b(\d{1,2})\s*(years?|yrs?|ans?|months?|mois)\b"
    r"|\b(\d{1,2})\s*(years?|yrs?|ans?|months?|mois)\s*(?:from|after|à compter de|après)\s*(?:the\s*)?"
    r"(?:date\s*of\s*|la date de\s*)?(?:manufactur|fabrication|production)", re.I)


def _duree_validite(documents) -> tuple[str, object] | None:
    """« Shelf life: 5 years », « 3 years from date of manufacture » -> « 5 ans ».
    Pas la déclaration de conformité : sa « validity » est celle du certificat."""
    from collections import Counter
    for code in ("etiquetage", "notice", "catalogue"):
        for piece, texte in _textes(documents, code):
            trouves = Counter()
            for m in DUREE.finditer(re.sub(r"\s+", " ", texte)):
                nombre, unite = (m.group(1), m.group(2)) if m.group(1) else (m.group(3), m.group(4))
                trouves[f"{int(nombre)} {'mois' if unite.lower().startswith('m') else 'ans'}"] += 1
            if trouves:
                return trouves.most_common(1)[0][0], piece
    # étiquette « MFG 2020-11-09 … EXP 2023-11-08 » : durée entre fabrication et péremption
    for code in ("etiquetage", "notice", "catalogue"):
        for piece, texte in _textes(documents, code):
            plat = re.sub(r"\s+", " ", texte)
            fab = re.search(r"(?:\bMFG\b|\bMFD\b|Mfg\.? ?Date|manufactur\w* date|date of manufacture|production date|"
                            r"date de fabrication)\s*[:.]?\s*" + DATE, plat, re.I)
            per = re.search(r"(?:\bEXP\b|expiry date|expiration date|use[- ]by|date de p[ée]remption)\s*[:.]?\s*" + DATE,
                            plat, re.I)
            if fab and per and (duree := duree_entre(_dates(fab.group(0))[0], _dates(per.group(0))[0])):
                return duree, piece
    return None


def _presentation(documents) -> tuple[str, object] | None:
    """Dispositif stérile à usage unique : « Emballage unitaire stérile », avec
    la boîte lue sur l'étiquette (« Qty: 10 » -> « /Boite de 10 unités stériles »),
    comme dans les dossiers acceptés."""
    if not _sterilisation(documents):
        return None
    for code in ("etiquetage", "catalogue", "notice"):
        for piece, texte in _textes(documents, code):
            plat = re.sub(r"\s+", " ", texte)
            if not re.search(r"single[- ]use|disposable|do not re-?use|usage unique|jetable|\(2\)|sterile", plat, re.I):
                continue
            m = (re.search(r"\b(?:Qty|Quantity|Quantit[ée])\s*[:.]?\s*(\d{1,4})\b", plat, re.I)
                 or re.search(r"\b(\d{1,4})\s*(?:pcs|pieces|units|unit[ée]s|pairs|paires)\s*(?:/|per|par)\s*(?:box|bo[iî]te)",
                              plat, re.I)
                 or re.search(r"\b(?:box|bo[iî]te) (?:of|de) (\d{1,4})\b", plat, re.I))
            n = int(m.group(1)) if m else 0
            return "Emballage unitaire stérile" + (f"/Boite de {n} unités stériles" if n > 1 else ""), piece
    return None


def _organisme_notifie(documents, ce_organisme: dict) -> str | None:
    """N° à 4 chiffres de l'organisme notifié : lu sur le certificat CE, sinon le
    « CE 0123 » le plus fréquent des étiquettes, de la déclaration ou de la notice."""
    from collections import Counter
    if (m := re.search(r"\b(\d{4})\b", ce_organisme.get("valeur") or "")):
        return m.group(1)
    trouves = Counter()
    for code in ("etiquetage", "declaration_conformite", "notice", "piece_specifique"):
        for _, texte in _textes(documents, code):
            trouves.update(re.findall(r"\bCE\s?(\d{4})\b", texte))
    return trouves.most_common(1)[0][0] if trouves else None


def _courriel(documents) -> tuple[str, object] | None:
    """Adresse électronique la plus citée dans les documents du fabricant."""
    from collections import Counter
    for code in ("declaration_conformite", "etiquetage", "notice", "attestation_fabricant", "catalogue"):
        for piece, texte in _textes(documents, code):
            trouves = Counter(m.lower() for m in re.findall(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", texte))
            if trouves:
                return trouves.most_common(1)[0][0], piece
    return None


RELEVES = {"sterilisation": _sterilisation, "duree_validite": _duree_validite, "presentation": _presentation}


def _transition_2023_607(dossier, une) -> tuple[str, str] | None:
    """Certificat MDD prolongé par une lettre au titre du Règlement (UE) 2023/607 :
    date de fin de la période de transition (31/12/2027 pour une classe III ou
    un IIb implantable hors « WET », 31/12/2028 pour les autres IIb, les IIa et
    les I), lue sur les lignes de la lettre qui citent le produit et sa classe.
    Plusieurs dates possibles : les deux sont indiquées, à préciser."""
    for piece, texte in _textes(dossier.documents, "piece_specifique"):
        if "2023/607" not in texte:
            continue
        noms = {m.lower() for m in re.findall(r"[A-Za-zÀ-ÿ]{4,}", " ".join(
            filter(None, [une("nom_marque")["valeur"], une("references")["valeur"]])))} - {"voir", "annexe"}
        classe_dossier = {"IIB": "iib", "IIA": "iia", "III": "iii"}.get((dossier.classe or "").upper(), "i")
        plat = re.sub(r"\s+", " ", texte)
        lignes = []
        for m in re.finditer(r"Class (III|IIb|IIa|I)\b", plat):
            if plat[max(0, m.start() - 10):m.start()].lower().endswith("excluding "):
                continue  # « Class IIb excluding Class IIb implantable » : une seule ligne
            # nom du dispositif : ce qui précède la classe depuis la fin de la ligne précédente du tableau
            avant = re.split(r"\)|N/A", plat[max(0, m.start() - 90):m.start()])[-1]
            avant = re.sub(r"B-UDI \S+", "", avant).strip()
            if m.group(1).lower() != classe_dossier or not any(n in avant.lower() for n in noms):
                continue
            classe = plat[m.start():m.start() + 60].split(" N/A")[0].strip()
            if "excluding" in classe.lower() or m.group(1) in ("IIa", "I"):
                date = "31/12/2028"
            elif m.group(1) == "III" or "implantable" in classe.lower():
                date = "31/12/2027"
            else:
                date = "31/12/2028"
            score = sum(n in avant.lower() for n in noms)
            lignes.append((score, date, f"{avant} : {classe} → {date}"))
        meilleur = max((l[0] for l in lignes), default=0)
        dates = {l[1] for l in lignes if l[0] == meilleur}  # lignes qui citent le plus de mots du produit
        lignes = [l[2] for l in lignes if l[0] == meilleur]
        if len(dates) == 1:
            return (f"jusqu'au {dates.pop()} (période de transition, Règlement (UE) 2023/607)",
                    f"Lettre 2023/607 (pièce {piece.numero}) — {lignes[0]}")
        if dates:
            return ("jusqu'au 31/12/2027 ou 31/12/2028 (période de transition, Règlement (UE) 2023/607) — à préciser",
                    f"Lettre 2023/607 (pièce {piece.numero}), produits de la gamme aux deux dates — "
                    + " ; ".join(dict.fromkeys(lignes)))
        if classe_dossier in ("iia", "i"):
            return ("jusqu'au 31/12/2028 (période de transition, Règlement (UE) 2023/607)",
                    f"Règlement (UE) 2023/607, classe {dossier.classe} (lettre en pièce {piece.numero})")
        return ("jusqu'au 31/12/2027 ou 31/12/2028 (période de transition, Règlement (UE) 2023/607) — à préciser",
                f"Règlement (UE) 2023/607 : 31/12/2027 pour un IIb implantable ou un III, sinon 31/12/2028 "
                f"(lettre en pièce {piece.numero})")
    return None


SOCIETE = re.compile(r"\b(?:medical|devices?|ltd|limited|pvt|private|co|corp|inc|s\.?p\.?a|gmbh|sarl|technology|"
                     r"industr\w*|group|company|rubbers?)\b", re.I)


def _marque_courte(marque: str, fabricant: str | None) -> str:
    """« ILIFE MEDICAL DEVICES » (raison sociale) -> « ILIFE » : la marque, comme
    dans les dossiers acceptés (« iLife / OneFlon® … »)."""
    mots = marque.split()
    if len(mots) > 1 and (SOCIETE.search(marque) or (fabricant and marque.lower() in fabricant.lower())):
        return mots[0].strip(",.")
    return marque


def _plusieurs_produits(resultat: dict) -> None:
    """Produits lus (catalogue, étiquettes) : une ligne par produit en annexe,
    avec la marque ; dans le formulaire, « MARQUE / Nom commercial » quand il
    n'y a qu'un nom commercial, « MARQUE / Voir annexe » sinon, comme dans les
    dossiers acceptés."""
    references, nom = resultat.get("references"), resultat.get("nom_marque")
    if not references or not references["valeur"] or references["provenance"] in ("saisie", "memoire"):
        return
    fabricant = (resultat.get("fabricant_nom") or {}).get("valeur")
    marque = ""
    if nom and nom["valeur"] and nom["provenance"] != "saisie":
        marque = _marque_courte(re.split(r"\s*/\s*", nom["valeur"])[0].strip(), fabricant).upper()
    elif fabricant:  # marque = nom du fabricant
        marque = fabricant.split()[0].strip(",.").upper()
    # « WZDSS-A-35W [LOT 2409015 Qty:10 » : le lot et la quantité ne sont pas des références
    propre = re.sub(r"\[?\s*\b(?:LOT|Qty|Quantity)\b\s*[:.]?\s*[\w-]*\]?", "", references["valeur"], flags=re.I)
    propre = re.sub(r"\bREF\s*[:.]\s*", "", propre)
    lignes = list(dict.fromkeys(re.sub(r"\s*\|\s*$", "", re.sub(r"\s+", " ", l)).strip()
                                for l in propre.splitlines() if l.strip()))
    if lignes and all(l.count("|") == 2 for l in lignes):  # déjà « MARQUE | NOM | RÉF »
        produits = [[p.strip() for p in l.split("|")][1:] for l in lignes]
        if marque:
            lignes = [f"{marque} | {p[0]} | {p[1]}" for p in produits]
        references["valeur"] = "\n".join(lignes)
    else:
        if any("|" in l for l in lignes):  # « NOM | REF » (étiquettes) : la marque est ajoutée
            produits = [[p.strip() for p in l.split("|")][:2] for l in lignes]
        else:
            produits = [[p.strip(), ""] for p in re.split(r"[,;\n]", "\n".join(lignes)) if p.strip()]
        references["valeur"] = "\n".join(f"{marque} | {p[0]} | {p[1] if len(p) > 1 else ''}" for p in produits)
    if not (produits and marque and nom and nom["provenance"] != "saisie"):
        return
    if "annexe" in (nom["valeur"] or "").lower() or "/" in (nom["valeur"] or ""):
        return
    if not nom["valeur"]:  # marque tirée du nom du fabricant : à vérifier
        nom.update(provenance="donnee", a_verifier=True, detail="Marque = nom du fabricant")
    noms = list(dict.fromkeys(p[0].strip() for p in produits if p[0].strip()))
    if len(noms) == 1:  # un seul nom commercial : « MARQUE / Nom » (dossiers acceptés)
        nom["valeur"] = f"{marque} / {noms[0]}"
        nom["detail"] = (nom["detail"] or "") + (f" — {len(produits)} références : détail en annexe"
                                                 if len(produits) > 1 else " — nom commercial lu dans les références")
    elif len(produits) > 1:
        nom["valeur"] = f"{marque} / Voir annexe"
        nom["detail"] = (nom["detail"] or "") + f" — {len(produits)} produits : détail en annexe"


# ---------------------------------------------------------------- DOCX


def _lignes_references(valeur: str | None) -> list[list[str]]:
    """« MARQUE | NOM | REF » par ligne ; une ligne sans séparateur va dans la
    colonne « Noms commerciaux »."""
    lignes = []
    for brute in (valeur or "").splitlines():
        if not brute.strip():
            continue
        parties = [p.strip() for p in brute.replace("\t", "|").split("|")]
        lignes.append((parties + ["", "", ""])[:3] if len(parties) > 1 else ["", brute.strip(), ""])
    return lignes


def _ecrire_valeur(cellule, donnee: dict) -> None:
    from docx.enum.text import WD_COLOR_INDEX

    paragraphe = cellule.paragraphs[0]
    valeur = donnee["valeur"]
    if donnee["type"] == "choix" and donnee["options"]:
        autres = valeur and valeur not in donnee["options"]
        texte = "   ".join(f"{'☒' if o == valeur else '☐'} {o}" for o in donnee["options"])
        paragraphe.add_run(texte + (f"   ☒ {valeur}" if autres else ""))
        if not valeur:
            paragraphe.add_run(f"   {A_COMPLETER}").font.highlight_color = WD_COLOR_INDEX.YELLOW
        return
    if not valeur:
        paragraphe.add_run(A_COMPLETER).font.highlight_color = WD_COLOR_INDEX.YELLOW
        return
    premiere, *suite = valeur.splitlines() or [""]
    paragraphe.add_run(premiere)
    for ligne in suite:
        cellule.add_paragraph(ligne)


def remplir_docx(code: str, donnees: dict[str, dict], titre_piece: str, chemin: Path) -> dict:
    """Écrit le formulaire `code` rempli ; retourne le bilan (cases vides, à vérifier)."""
    from docx import Document
    from docx.shared import Pt

    config = charger()["formulaires"][code]
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(titre_piece, level=1)
    doc.add_paragraph().add_run(MENTION_VALIDATION).bold = True
    doc.add_heading(config["titre"], level=2)
    utilisees = []
    for section in config["sections"]:
        doc.add_heading(section["titre"], level=3)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for ligne in section["lignes"]:
            ident, libelle = (next(iter(ligne.items())) if isinstance(ligne, dict) else (ligne, None))
            donnee = donnees[ident]
            utilisees.append(donnee)
            cellules = table.add_row().cells
            cellules[0].paragraphs[0].add_run(libelle or donnee["libelle"]).bold = True
            _ecrire_valeur(cellules[1], donnee)
    doc.add_paragraph()
    representant = donnees["representant_legal"]["valeur"] or A_COMPLETER
    doc.add_paragraph(config["signature"].format(representant_legal=representant))
    if config.get("partie_reservee"):
        doc.add_paragraph().add_run(config["partie_reservee"]).italic = True
    if config.get("annexe"):
        doc.add_page_break()
        titre, *suite = config["annexe"].splitlines()
        doc.add_heading(titre, level=2)
        for ligne in suite:
            doc.add_paragraph(ligne)
        references = donnees["references"]
        utilisees.append(references)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cellule, entete in zip(table.rows[0].cells, ("Nom de marque", "Noms commerciaux", "Référence")):
            cellule.paragraphs[0].add_run(entete).bold = True
        lignes = _lignes_references(references["valeur"]) or [[A_COMPLETER, "", ""]]
        for ligne in lignes:
            for cellule, texte in zip(table.add_row().cells, ligne):
                cellule.paragraphs[0].add_run(texte)

    vides = [d for d in utilisees if not d["valeur"]]
    a_verifier = [d for d in utilisees if d["valeur"] and d["a_verifier"]]
    doc.add_heading("Origine des données (relecture — à retirer avant dépôt)", level=2)
    doc.add_paragraph(f"{len(utilisees) - len(vides)} case(s) remplie(s), {len(vides)} à compléter, "
                      f"{len(a_verifier)} à vérifier.")
    for d in utilisees:
        if d["valeur"] and d["provenance"] not in ("saisie",):
            marque = " — À VÉRIFIER" if d["a_verifier"] else ""
            doc.add_paragraph(f"{d['libelle']} : {d['detail']}{marque}", style="List Bullet")
    chemin.parent.mkdir(parents=True, exist_ok=True)
    doc.save(chemin)
    return {"remplies": len(utilisees) - len(vides), "a_completer": len(vides), "a_verifier": len(a_verifier)}


def generer_formulaire(dossier, code: str, titre_piece: str, dossier_sortie: Path, profil: dict) -> dict:
    """Remplit le formulaire d'une pièce ; même forme de résultat que generate.generer_document."""
    donnees = resoudre(dossier, profil)
    chemin = dossier_sortie / f"{code}.docx"
    bilan = remplir_docx(code, donnees, titre_piece, chemin)
    regles = charger()
    return {
        "fichier": str(chemin),
        "sources": [{
            "texte_source": (f"Formulaire officiel rempli par le code (rules/formulaires_maroc.yaml v{regles['version']}) : "
                             f"{bilan['remplies']} case(s) remplie(s), {bilan['a_completer']} à compléter, "
                             f"{bilan['a_verifier']} à vérifier"),
            "date_version": regles["date_version"],
            "fichier": "rules/formulaires_maroc.yaml",
        }],
        "bilan": bilan,
    }
