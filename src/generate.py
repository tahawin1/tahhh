"""
Pipeline de génération documentaire : Rule Engine → RAG → Mistral → DOCX.

Le LLM ne décide jamais des documents requis (c'est rule_engine.py qui
décide) — il ne fait que rédiger le contenu de chaque document en
s'appuyant sur des extraits réels retrouvés dans la base vectorielle
(dossiers déjà réalisés + textes réglementaires).

Usage :
    python3 src/generate.py --pays-origine chine --produit "Prothèse orthopédique de hanche" --classe IIB
"""
from __future__ import annotations

import argparse
import datetime
import os
import re
import time
import unicodedata
from pathlib import Path

import requests
from docx import Document
from qdrant_client import QdrantClient
from qdrant_client.models import FieldCondition, Filter, MatchValue

from embeddings import OLLAMA_BASE_URL, Embedder
from rule_engine import DocumentRequis, charger_regles, documents_requis_maroc, prochain_creneau_depot

COLLECTION = "dossiers_reference"
OLLAMA_URL = f"{OLLAMA_BASE_URL}/api/generate"
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "mistral")
# Sur CPU seul, Mistral 7B rédige ~5 tokens/s : prévoir large.
OLLAMA_TIMEOUT = int(os.environ.get("OLLAMA_TIMEOUT", "900"))
# Longueur maximale de chaque document généré (tokens) ; -1 = sans limite.
OLLAMA_NUM_PREDICT = int(os.environ.get("OLLAMA_NUM_PREDICT", "700"))
OUTPUT_DIR = Path(__file__).resolve().parent.parent / "output"
PROFIL_ENTREPRISE = Path(__file__).resolve().parent.parent / "config" / "entreprise.yaml"
LIBELLES_PROFIL = {
    "raison_sociale": "Raison sociale", "ville": "Ville (lieu des courriers)", "adresse": "Adresse",
    "telephone": "Téléphone", "email": "Adresse électronique", "representant_legal": "Représentant légal", "ice": "ICE",
    "identifiant_fiscal": "IF", "registre_commerce": "RC", "patente": "Patente", "banque_rib": "Banque / RIB",
}


def profil_entreprise() -> dict:
    """Coordonnées de l'établissement demandeur (config/entreprise.yaml, sur le
    serveur uniquement) : reprises telles quelles, jamais inventées."""
    if not PROFIL_ENTREPRISE.exists():
        return {}
    import yaml
    with open(PROFIL_ENTREPRISE, encoding="utf-8") as f:
        brut = {k: _reparer(str(v).strip()) for k, v in (yaml.safe_load(f) or {}).items() if k in LIBELLES_PROFIL and v}
    # une valeur restée égale à celle du fichier d'exemple (« Nom du représentant légal ») n'est pas une donnée
    exemple = {}
    if (fichier_exemple := PROFIL_ENTREPRISE.with_name("entreprise.exemple.yaml")).exists():
        with open(fichier_exemple, encoding="utf-8") as f:
            exemple = {k: normaliser_libelle(str(v)) for k, v in (yaml.safe_load(f) or {}).items() if v}
    return {k: v for k, v in brut.items()
            if v and normaliser_libelle(v) != exemple.get(k) and not v.lower().startswith("[à compléter")}


def normaliser_libelle(texte: str) -> str:
    return re.sub(r"\s+", " ", texte.strip().lower())


def _reparer(texte: str) -> str:
    """Accents abîmés par une console Windows (UTF-8 relu en CP850 ou CP1252 :
    « repr├®sentant », « reprÃ©sentant ») : rétablis, sinon texte inchangé."""
    if not any(c in texte for c in "├Ã"):
        return texte
    for codage in ("cp1252", "cp850"):
        try:
            repare = texte.encode(codage).decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            continue
        if not any(c in repare for c in "├Ã"):
            return repare
    return texte


# Destinataire des demandes, lu dans les règles (AMMPS depuis la v2) — jamais décidé par le LLM
DESTINATAIRE = charger_regles("maroc").get("destinataire_demande", "[À COMPLÉTER : destinataire]")

# Pays dont les textes sont indexés dans Qdrant (voir scripts/indexer_tout.sh)
PAYS_INDEXES = {"maroc", "chine", "inde", "union_europeenne", "etats_unis", "coree_du_sud", "pakistan"}

MENTION_VALIDATION = (
    "PROJET GÉNÉRÉ AUTOMATIQUEMENT — EN ATTENTE DE VALIDATION HUMAINE. "
    "Ne pas déposer avant relecture et validation explicite."
)


PREAMBULE = re.compile(r"^\s*(voici|voil[aà]|ci-dessous|ci-apr[eè]s)\b.*:\s*$", re.I)


def _mots_significatifs(ligne: str) -> set[str]:
    from extraction import normaliser
    return {m for m in re.findall(r"\w+", normaliser(ligne)) if len(m) > 2}


def recadrer_sur_modele(texte: str, modele: str) -> tuple[str, int]:
    """Garde-fou déterministe (le prompt seul ne suffit pas avec Mistral 7B) :
    retire une phrase d'introduction (« Voici la lettre… : ») et tout ce qui
    suit la fin du modèle (listes ou consignes inventées après la signature ou
    le pied de page). La fin du modèle est la dernière de ses lignes retrouvée
    dans le texte (le pied de page peut manquer) ; les lignes courtes qui la
    suivent (nom du signataire) sont gardées. Retourne (texte recadré, nombre
    de lignes retirées à la fin)."""
    lignes = texte.strip().splitlines()
    while lignes and (not lignes[0].strip() or PREAMBULE.match(lignes[0])):
        lignes.pop(0)
    reperes = [_mots_significatifs(l) for l in modele.strip().splitlines()]
    reperes = [r for r in reperes if len(r) >= 2]
    dernier = None
    for repere in reversed(reperes[-8:]):  # de la dernière ligne du modèle vers le haut
        seuil = len(repere) if len(repere) <= 3 else max(3, 0.6 * len(repere))
        trouves = [i for i, l in enumerate(lignes) if len(_mots_significatifs(l) & repere) >= seuil]
        if trouves:
            dernier = trouves[-1]
            break
    if dernier is None:
        return "\n".join(lignes).strip(), 0
    while (dernier + 1 < len(lignes) and len(lignes[dernier + 1].split()) <= 5
           and not lignes[dernier + 1].rstrip().endswith(":")):
        dernier += 1  # signataire, ligne vide — pas un titre recopié du prompt (« Consignes strictes : »)
    retirees = [l for l in lignes[dernier + 1:] if l.strip()]
    return "\n".join(lignes[: dernier + 1]).strip(), len(retirees)


LIEU_DATE = re.compile(r"^(\s*)([^,\n]{2,40}),\s*le\s+\S+(.*)$", re.I)


def dater(texte: str, ville: str | None, jour: datetime.date | None = None) -> str:
    """Ligne « Ville, le jj/mm/aaaa » : date du jour et ville du profil (Mistral
    recopie celles du modèle). Seules les premières lignes sont regardées."""
    jour = jour or datetime.date.today()
    lignes = texte.splitlines()
    for i, ligne in enumerate(lignes[:10]):
        if m := LIEU_DATE.match(ligne):
            lignes[i] = f"{m.group(1)}{ville or m.group(2).strip()}, le {jour:%d/%m/%Y}"
            break
    return "\n".join(lignes)


class GenerationImpossible(RuntimeError):
    """La génération d'un document n'a pas pu aboutir (ou n'est pas permise)."""


class GenerateurDocuments:
    def __init__(self, host: str | None = None, port: int | None = None):
        self.client = QdrantClient(
            host=host or os.environ.get("QDRANT_HOST", "localhost"),
            port=port or int(os.environ.get("QDRANT_PORT", "6333")),
        )
        self.embedder = Embedder()
        self._session_ollama = requests.Session()
        self._session_ollama.trust_env = False  # Ollama tourne en local, pas de proxy

    def rechercher_references(self, pays: str, requete: str, k: int = 3) -> list[dict]:
        """Recherche les chunks les plus pertinents pour un pays donné (RAG)."""
        vecteur = self.embedder.encoder_un(requete)
        resultats = self.client.query_points(
            collection_name=COLLECTION,
            query=vecteur,
            query_filter=Filter(must=[FieldCondition(key="pays", match=MatchValue(value=pays))]),
            limit=k,
            with_payload=True,
        ).points
        return [{**r.payload, "score": r.score} for r in resultats]

    def interroger_json(self, prompt: str, schema: dict) -> str:
        """Réponse de Mistral au format JSON imposé (sortie structurée d'Ollama),
        température 0 ; une erreur 5xx passagère est relancée une fois."""
        for tentative in (1, 2):
            r = self._session_ollama.post(OLLAMA_URL, json={
                "model": OLLAMA_MODEL, "prompt": prompt, "stream": False, "format": schema,
                "options": {"temperature": 0, "num_ctx": 8192, "num_predict": 400}}, timeout=OLLAMA_TIMEOUT)
            if r.status_code < 500 or tentative == 2:
                break
            time.sleep(5)
        r.raise_for_status()
        return r.json()["response"]

    def modele_pour(self, document: DocumentRequis, pays_origine: str, produit: str,
                    classe: str | None = None) -> dict | None:
        """Pièce équivalente d'un dossier accepté, ou None (rédaction sans modèle)."""
        from modeles import modele_pour
        return modele_pour(document.id, pays_origine, produit, self.client, self.embedder, classe=classe)

    def references_pour_document(self, document: DocumentRequis, pays_destination: str, pays_origine: str) -> list[dict]:
        """
        Choisit (par le code, pas par le LLM) dans quels textes chercher :
        toujours ceux du pays de destination ; en plus, pour la pièce
        spécifique à l'origine (certificat NMPA / CDSCO / CE), ceux du pays
        d'origine qui délivre ce certificat.
        """
        requete = f"{document.nom}. {document.consigne_redaction or ''}".strip()
        references = self.rechercher_references(pays_destination, requete, k=3)
        if document.origine_regle.startswith("piece_specifique") and pays_origine in PAYS_INDEXES:
            references += self.rechercher_references(pays_origine, document.nom, k=2)
        return references

    def generer_contenu(
        self,
        document: DocumentRequis,
        pays_destination: str,
        produit: str,
        pays_origine: str,
        classe: str | None = None,
        fournisseur: str | None = None,
        pieces_du_dossier: list[str] | None = None,
        donnees: dict[str, str] | None = None,
    ) -> tuple[str, list[dict]]:
        """Appelle Mistral (via Ollama) pour rédiger le contenu d'un document,
        en s'appuyant sur les extraits réglementaires retrouvés par RAG.
        Retourne (contenu, références utilisées) — les références sont
        remontées à la validation humaine pour vérification.

        Lève GenerationImpossible si le document n'est pas à rédiger ou si
        Mistral ne répond pas : jamais de document de remplacement."""
        if not document.a_rediger:
            raise GenerationImpossible(
                f"'{document.nom}' est une pièce à fournir par {document.fourni_par} : "
                "elle n'est jamais rédigée par le système."
            )

        # Avec un modèle accepté, c'est lui qui fixe la forme : les extraits de textes
        # (souvent d'autres procédures) ne sont pas donnés, pour ne rien y recopier.
        modele = self.modele_pour(document, pays_origine, produit, classe)
        references = [] if modele else self.references_pour_document(document, pays_destination, pays_origine)
        contexte = "\n\n".join(
            f"<<< EXTRAIT {i} — {r['texte_source']} (version du {r['date_version']})\n{r['texte'][:800]}\n>>>"
            for i, r in enumerate(references, 1)
        ) or "(aucun extrait trouvé dans la base — le signaler à la validation humaine)"
        sources = [
            {
                "texte_source": r["texte_source"],
                "pays": r["pays"],
                "date_version": r["date_version"],
                "fichier": r["fichier"],
                "chunk_index": r["chunk_index"],
                "score": round(r["score"], 4),
            }
            for r in references
        ]

        # Modèle : la même pièce dans un dossier réellement accepté (voir modeles.py)
        if modele:
            sources.insert(0, {
                "texte_source": f"Modèle : pièce {modele['numero']} du dossier accepté « {modele['produit']} »",
                "pays": modele["pays_origine"],
                "date_version": modele["indexe_le"],
                "fichier": modele["fichier"],
                "chunk_index": 0,
                "score": round(modele["score"], 4),
            })
            bloc_modele = f"""
MODÈLE À SUIVRE — la même pièce, dans un dossier réellement accepté par l'administration pour un AUTRE dispositif (« {modele['produit']} ») :
<<< MODÈLE
{modele['texte']}
>>>
Comment utiliser le modèle :
- Reprends sa structure, son ordre, son destinataire et ses formules. N'ajoute AUCUNE rubrique, liste ou paragraphe absent du modèle.
- Les coordonnées de l'établissement demandeur sont celles du profil ci-dessus (identiques à celles du modèle).
- Remplace TOUT ce qui décrit l'ancien dispositif « {modele['produit']} » (désignation, noms commerciaux, références, présentation, indications, intérêt médical, domaine thérapeutique, composition, fabricant, classe, codes, dates) par les données du nouveau dossier ; une donnée inconnue du nouveau dossier devient [À COMPLÉTER]. Ne laisse aucune donnée propre à l'ancien dispositif, et n'en déduis aucune pour le nouveau.
"""
        else:
            bloc_modele = ""

        profil = profil_entreprise()
        bloc_profil = (
            "Établissement demandeur (coordonnées à reprendre exactement) :\n"
            + "\n".join(f"  - {LIBELLES_PROFIL[k]} : {v}" for k, v in profil.items()) + "\n"
        ) if profil else "Établissement demandeur : [À COMPLÉTER] (profil config/entreprise.yaml absent)\n"

        # Cadre réglementaire du pays d'origine jusqu'au Maroc (rules/correspondances.yaml)
        try:
            from correspondances import fiche_pays
            bloc_cadre = ("Cadre réglementaire (fixé par les règles, à ne jamais contredire ; ne cite un article que s'il figure ici) :\n"
                          + fiche_pays(pays_origine, classe) + "\n")
        except Exception:  # fichier absent ou invalide : rédaction sans ce bloc
            bloc_cadre = ""
        # Données déjà établies (saisies par l'utilisateur ou lues dans les documents reçus)
        bloc_donnees = "".join(f"- {k} : {v}\n" for k, v in (donnees or {}).items())
        consigne = document.consigne_redaction or document.nom
        # La liste des pièces vient du moteur de règles, jamais du LLM
        liste_pieces = (
            "Pièces composant le dossier (liste fixée par le moteur de règles — ne rien ajouter ni retirer) :\n"
            + "\n".join(f"  {i}. {nom}" for i, nom in enumerate(pieces_du_dossier, 1))
            + "\n"
        ) if pieces_du_dossier else ""
        prompt = f"""Tu es un assistant spécialisé en constitution de dossiers réglementaires pour dispositifs médicaux au Maroc.

DOCUMENT À RÉDIGER : {document.nom}
Ce que ce document doit contenir : {consigne}

Informations connues sur le dossier (à reprendre telles quelles) :
- Dispositif médical : {produit}
- Classe du dispositif (classification marocaine) : {classe or "[À COMPLÉTER]"}
- Fabricant / fournisseur : {fournisseur or "[À COMPLÉTER]"}
- Pays d'origine du fournisseur : {pays_origine}
- Pays de destination du dossier : {pays_destination}
{bloc_donnees}{bloc_profil}{liste_pieces}{bloc_cadre}{bloc_modele}
{"" if modele else "Extraits de textes réglementaires officiels, fournis UNIQUEMENT comme référence (pour les exigences et le vocabulaire) :" + chr(10) + contexte}

Consignes strictes :
- Rédige le document demandé lui-même, prêt à être complété et signé. Ne recopie PAS les extraits, ne reproduis pas d'en-têtes du Bulletin officiel, ne cite pas les extraits dans le document.
- N'invente aucune information factuelle (nom, adresse, numéro, date, référence) : écris [À COMPLÉTER] à la place.
- Intérêt médical, indications, références, composition : uniquement ce que fournit le nouveau dossier, sinon [À COMPLÉTER] — ne les déduis jamais du nom du dispositif.
- N'aborde que ce document, sans parler d'autres procédures (publicité, inspection, sanctions…).
- Si le document énumère les pièces du dossier, reprends EXACTEMENT la liste fixée ci-dessus, sans en ajouter ni en retirer.
- Destinataire des demandes : {DESTINATAIRE} (jamais le ministre ni la DMP).
- Rédige en français, dans le registre administratif marocain (formule de politesse administrative, aucune formule familière), sans commentaire avant ou après le document."""

        try:
            # Une erreur 500 passagère d'Ollama (génération interrompue) est
            # relancée une fois ; à la seconde, la pièce passe en erreur.
            for tentative in (1, 2):
                reponse = self._session_ollama.post(
                    OLLAMA_URL,
                    json={
                        "model": OLLAMA_MODEL,
                        "prompt": prompt,
                        "stream": False,
                        # contexte élargi : le modèle + les extraits dépassent les 4096 jetons par défaut ;
                        # un long modèle (fiche signalétique) demande une réponse plus longue
                        "options": {"num_ctx": 8192,
                                    "num_predict": max(OLLAMA_NUM_PREDICT, len(modele["texte"]) // 3) if modele else OLLAMA_NUM_PREDICT},
                    },
                    timeout=OLLAMA_TIMEOUT,
                )
                if reponse.status_code < 500 or tentative == 2:
                    break
                time.sleep(5)
            if reponse.status_code >= 500:
                raise GenerationImpossible(
                    f"Mistral a échoué deux fois (erreur {reponse.status_code} d'Ollama) — "
                    "voir `journalctl -u ollama`, puis relancer la rédaction."
                )
            reponse.raise_for_status()
            texte = reponse.json()["response"]
            if modele:
                texte, retirees = recadrer_sur_modele(texte, modele["texte"])
                if retirees:
                    sources[0]["recadrage"] = f"{retirees} ligne(s) ajoutée(s) par Mistral après la fin du modèle, retirée(s)"
                texte = dater(texte, profil_entreprise().get("ville"))
            return texte, sources
        except requests.exceptions.ConnectionError as e:
            raise GenerationImpossible(
                f"Ollama n'est pas démarré sur {OLLAMA_URL} — lancer `ollama serve` puis `ollama pull {OLLAMA_MODEL}`."
            ) from e
        except requests.exceptions.Timeout as e:
            raise GenerationImpossible(
                f"Mistral n'a pas répondu en {OLLAMA_TIMEOUT}s — augmenter OLLAMA_TIMEOUT ou utiliser un GPU."
            ) from e

    def creer_fichier_docx(
        self,
        titre: str,
        contenu: str,
        chemin_sortie: Path,
        document: DocumentRequis | None = None,
        sources: list[dict] | None = None,
    ) -> Path:
        doc = Document()
        doc.add_heading(titre, level=1)
        doc.add_paragraph().add_run(MENTION_VALIDATION).bold = True
        if document is not None:
            if document.traduction_requise:
                doc.add_paragraph("⚠ Traduction assermentée requise pour la pièce justificative.")
            if document.legalisation_requise:
                doc.add_paragraph("⚠ Légalisation / apostille requise pour la pièce justificative.")
        for paragraphe in contenu.split("\n"):
            if paragraphe.strip():
                doc.add_paragraph(paragraphe)
        if sources:
            doc.add_heading("Sources réglementaires utilisées (à vérifier)", level=2)
            for s in sources:
                doc.add_paragraph(
                    f"{s['texte_source']} — version du {s['date_version']} "
                    f"({s['fichier']}, extrait n°{s['chunk_index']})",
                    style="List Bullet",
                )
        chemin_sortie.parent.mkdir(parents=True, exist_ok=True)
        doc.save(chemin_sortie)
        return chemin_sortie


def nouveau_dossier_sortie(pays_origine: str, produit: str) -> Path:
    """Un dossier de sortie par dossier : deux dossiers ne s'écrasent jamais."""
    horodatage = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    return OUTPUT_DIR / f"{horodatage}_{pays_origine}_{_slug(produit)}"


def generer_document(
    generateur: GenerateurDocuments,
    doc: DocumentRequis,
    produit: str,
    pays_origine: str,
    dossier_sortie: Path,
    classe: str | None = None,
    fournisseur: str | None = None,
    pieces_du_dossier: list[str] | None = None,
    donnees: dict[str, str] | None = None,
) -> dict:
    """Rédige UN document à rédiger (RAG + Mistral) et l'enregistre en DOCX.
    Lève GenerationImpossible pour une pièce à fournir ou si Mistral échoue."""
    contenu, sources = generateur.generer_contenu(
        doc, pays_destination="maroc", produit=produit, pays_origine=pays_origine,
        classe=classe, fournisseur=fournisseur, pieces_du_dossier=pieces_du_dossier, donnees=donnees,
    )
    chemin = generateur.creer_fichier_docx(
        doc.nom, contenu, dossier_sortie / f"{doc.id}.docx".replace(" ", "_"), document=doc, sources=sources
    )
    return {"fichier": str(chemin), "sources": sources}


def traiter_dossier(pays_origine: str, produit: str, classe: str | None = None) -> dict:
    """
    Fonction principale : orchestre rule_engine (décision) + RAG + Mistral
    (rédaction) + mise en forme DOCX. Retourne un résumé du dossier généré,
    prêt pour la validation humaine — jamais pour un dépôt automatique.

    Seules les pièces `a_rediger` sont rédigées ; les pièces `a_fournir`
    sont listées comme à obtenir auprès de leur émetteur.
    """
    documents = documents_requis_maroc(pays_origine, classe=classe)
    generateur = GenerateurDocuments()
    dossier_sortie = nouveau_dossier_sortie(pays_origine, produit)

    resultats = []
    for doc in documents:
        entree = {
            "id": doc.id,
            "document": doc.nom,
            "nature": doc.nature,
            "fourni_par": doc.fourni_par,
            "traduction_requise": doc.traduction_requise,
            "legalisation_requise": doc.legalisation_requise,
            "origine_regle": doc.origine_regle,
            "fichier": None,
            "sources": [],
            "erreur": None,
        }
        if doc.a_rediger:
            try:
                entree.update(generer_document(
                    generateur, doc, produit, pays_origine, dossier_sortie,
                    classe=classe, pieces_du_dossier=[d.nom for d in documents],
                ))
            except GenerationImpossible as e:
                entree["erreur"] = str(e)
        resultats.append(entree)

    return {
        "produit": produit,
        "pays_origine": pays_origine,
        "pays_destination": "maroc",
        "classe": classe,
        "dossier_sortie": str(dossier_sortie),
        "nombre_documents": len(resultats),
        "documents": resultats,
        "prochain_creneau_depot": str(prochain_creneau_depot()),
        "statut": "EN ATTENTE DE VALIDATION HUMAINE — aucun dépôt automatique",
    }


def _slug(texte: str) -> str:
    """'Prothèse orthopédique de hanche' -> 'prothese-orthopedique-de-hanche'"""
    ascii_ = unicodedata.normalize("NFKD", texte).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", ascii_.lower()).strip("-")[:60] or "dossier"


def main():
    parser = argparse.ArgumentParser(description="Génère un dossier de conformité complet")
    parser.add_argument("--pays-origine", required=True, choices=["chine", "inde", "union_europeenne", "etats_unis", "coree_du_sud", "pakistan", "autre"])
    parser.add_argument("--produit", required=True)
    parser.add_argument("--classe", default=None, choices=["I", "IS", "IM", "IR", "IIA", "IIB", "III"])
    args = parser.parse_args()

    resultat = traiter_dossier(args.pays_origine, args.produit, args.classe)

    print(f"\nDossier généré : {resultat['produit']} (origine : {resultat['pays_origine']})")
    print(f"Nombre de documents : {resultat['nombre_documents']}")
    for d in resultat["documents"]:
        flags = []
        if d["traduction_requise"]:
            flags.append("traduction requise")
        if d["legalisation_requise"]:
            flags.append("légalisation requise")
        suffix = f" [{', '.join(flags)}]" if flags else ""
        if d["nature"] == "a_fournir":
            cible = f"à obtenir auprès de {d['fourni_par']}"
        elif d["erreur"]:
            cible = f"ÉCHEC : {d['erreur']}"
        else:
            cible = d["fichier"]
        print(f"  - {d['document']}{suffix}\n    -> {cible}")
    print(f"\nProchain créneau de dépôt : {resultat['prochain_creneau_depot']}")
    print(f"Statut : {resultat['statut']}")


if __name__ == "__main__":
    main()
